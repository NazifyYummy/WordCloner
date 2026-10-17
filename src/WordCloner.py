"""
================================================================================
 WordCloner — Letter Issuance System
================================================================================
"""

import argparse
import os
import subprocess
import sys
import tempfile
from datetime import date, datetime
from json import load, dump
from shutil import copy2
from copy import deepcopy
from pathlib import Path
from collections import Counter

from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QLabel, QPushButton, QFrame, QGraphicsDropShadowEffect,
    QStackedWidget, QLineEdit, QFileDialog, QDialog, QComboBox,
    QScrollArea, QColorDialog, QCheckBox, QSystemTrayIcon, QMenu
)
from PySide6.QtCore import (
    Qt, QPoint, Signal, QSize, QRectF, QEvent, QTimer, QRegularExpression, QUrl
)
from PySide6.QtGui import (
    QIcon, QPainter, QColor, QFont, QPen, QPainterPath,
    QGuiApplication, QPalette, QRegularExpressionValidator, QAction,
    QDesktopServices,
)

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph as DocxParagraph
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False
    print("[warning] python-docx is not installed → pip install python-docx")

try:
    from openpyxl import load_workbook
    HAS_EXCEL = True
except ImportError:
    load_workbook = None
    HAS_EXCEL = False


# ═════════════════════════════════════════════════════════════════
#  Paths
# ═════════════════════════════════════════════════════════════════
from wordcloner_paths import (
    APP_ROOT, DOCUMENTS_DIR, WORDCLONER_DIR, WORDS_DIR,
    TEMPLATES_JSON, SETTINGS_JSON, migrate_legacy_data,
)

# ═════════════════════════════════════════════════════════════════

#  Settings
# ═════════════════════════════════════════════════════════════════

DEFAULT_SETTINGS = {
    "output_dir": str(Path.home() / "Desktop"),
    "placeholder_wrapper": '##',
    "quick_access": [],
    "language": "fa",     # fa | en
    "window_size": [720, 600],
    "window_maximized": False,
    "launch_on_startup": False,
}
SETTINGS = dict(DEFAULT_SETTINGS)


def load_settings():
    global SETTINGS
    SETTINGS = dict(DEFAULT_SETTINGS)
    if SETTINGS_JSON.exists():
        try:
            with open(SETTINGS_JSON, "r", encoding="utf-8") as f:
                data = load(f)
            if isinstance(data, dict):
                SETTINGS.update(data)
                if not isinstance(SETTINGS.get("quick_access"), list):
                    SETTINGS["quick_access"] = []
                if SETTINGS.get("language") not in ("fa", "en"):
                    SETTINGS["language"] = "fa"
                if not isinstance(SETTINGS.get("launch_on_startup"), bool):
                    SETTINGS["launch_on_startup"] = False
        except Exception as e:
            print(f"[error] reading settings: {e}")


def save_settings():
    WORDCLONER_DIR.mkdir(parents=True, exist_ok=True)
    try:
        with open(SETTINGS_JSON, "w", encoding="utf-8") as f:
            dump(SETTINGS, f, ensure_ascii=False, indent=4)
    except Exception as e:
        print(f"[error] saving settings: {e}")


STARTUP_RUN_KEY = r"Software\Microsoft\Windows\CurrentVersion\Run"
STARTUP_VALUE_NAME = "WordCloner"


def _windows_startup_command():
    """Return the command Windows should run at sign-in, without showing GUI."""
    executable = Path(sys.executable).resolve()
    command_parts = [str(executable)]

    if not getattr(sys, "frozen", False):
        # Prefer pythonw.exe for source runs so the startup entry doesn't flash
        # a console window. Fall back to the active interpreter if unavailable.
        if executable.name.lower() == "python.exe":
            pythonw = executable.with_name("pythonw.exe")
            if pythonw.is_file():
                command_parts[0] = str(pythonw)
        command_parts.append(str(Path(__file__).resolve()))

    command_parts.append("--startup")
    return subprocess.list2cmdline(command_parts)


def _set_windows_startup(enabled):
    """Create/remove the current user's Windows Run entry."""
    if sys.platform != "win32":
        raise OSError("Automatic startup is supported only on Windows.")

    import winreg

    key = winreg.CreateKeyEx(
        winreg.HKEY_CURRENT_USER,
        STARTUP_RUN_KEY,
        0,
        winreg.KEY_SET_VALUE,
    )
    try:
        if enabled:
            winreg.SetValueEx(
                key,
                STARTUP_VALUE_NAME,
                0,
                winreg.REG_SZ,
                _windows_startup_command(),
            )
        else:
            try:
                winreg.DeleteValue(key, STARTUP_VALUE_NAME)
            except FileNotFoundError:
                pass
    finally:
        winreg.CloseKey(key)


def _sync_windows_startup_from_settings():
    """Repair a missing/stale Run entry without blocking normal GUI startup."""
    if sys.platform != "win32":
        return
    try:
        _set_windows_startup(SETTINGS.get("launch_on_startup") is True)
    except Exception as exc:
        print(f"[startup] Could not synchronize Windows startup setting: {exc}")


# ═════════════════════════════════════════════════════════════════
#  i18n
# ═════════════════════════════════════════════════════════════════
from wordcloner_translations import T

LANG_MAP = {"fa": "persian", "en": "english"}


def tr(item, **kwargs):
    if not isinstance(item, dict):
        return str(item)
    key = LANG_MAP.get(SETTINGS.get("language", "fa"), "persian")
    value = item.get(key)
    if value is None:
        value = item.get("persian", "")
    if kwargs:
        try:
            value = value.format(**kwargs)
        except Exception:
            pass
    return value


def _is_rtl():
    return SETTINGS.get("language", "fa") == "fa"


def _dir():
    return Qt.RightToLeft if _is_rtl() else Qt.LeftToRight

#  Colors (only green theme)
# ═════════════════════════════════════════════════════════════════

C_DARK_GREEN  = "#2E5A4C"
C_GREEN       = "#3A7A5E"
C_LIGHT_GREEN = "#E8F0EC"
C_HOVER       = "#F0F5F2"
C_BG          = "#F4F7F2"
C_BORDER      = "#D0DDD5"
C_MUTED       = "#7A8F86"
C_ACCENT1     = "#8FB5A2"
C_ACCENT2     = "#E6EEE8"
C_RED         = "#E0453A"


# ═════════════════════════════════════════════════════════════════
#  Icons
# ═════════════════════════════════════════════════════════════════
from wordcloner_icons import (
    icon_document, icon_document_big, icon_plus_circle, icon_plus_white,
    icon_trash, icon_chevron_down, icon_grip, icon_upload, icon_arrow_left,
    icon_gear, icon_help, icon_status, icon_minimize, icon_maximize, icon_close,
    icon_check,
)


def apply_color_button_style(btn, hex_color):
    btn.setStyleSheet(f"""
        QPushButton {{
            background-color: {hex_color};
            border: 2px solid {C_GREEN};
            border-radius: 8px;
        }}
        QPushButton:hover {{
            border: 2px solid {C_DARK_GREEN};
        }}
    """)

# ═════════════════════════════════════════════════════════════════

#  ComboWithArrow
# ═════════════════════════════════════════════════════════════════

class ComboWithArrow(QComboBox):
    ARROW_AREA = 24

    _ARROW_QSS = """
        QComboBox::drop-down {
            border: none;
            width: 0px;
            background: transparent;
        }
        QComboBox::down-arrow {
            image: none;
            width: 0px;
            height: 0px;
        }
    """

    def __init__(self, editable=False, parent=None):
        super().__init__(parent)
        self.setEditable(editable)
        self.setLayoutDirection(_dir())
        self.setCursor(Qt.PointingHandCursor)

        if editable:
            le = self.lineEdit()
            if le is not None:
                le.setStyleSheet(
                    "QLineEdit { background: transparent; border: none; "
                    "padding: 0px; }"
                )
                self._update_text_margins()
                le.installEventFilter(self)

    def _is_rtl(self):
        return self.layoutDirection() == Qt.RightToLeft

    def _update_text_margins(self):
        le = self.lineEdit()
        if le is None:
            return
        if self._is_rtl():
            le.setTextMargins(self.ARROW_AREA, 0, 8, 0)
        else:
            le.setTextMargins(8, 0, self.ARROW_AREA, 0)

    def setStyleSheet(self, qss):
        qss = (qss or "")
        if not self._is_rtl():
            qss = qss.replace("padding-left: 26px;", "__PL__;")
            qss = qss.replace("padding-right: 10px;", "padding-right: 26px;")
            qss = qss.replace("__PL__;", "padding-left: 10px;")
        super().setStyleSheet(qss + self._ARROW_QSS)

    def eventFilter(self, obj, event):
        le = self.lineEdit()
        if obj is le and event.type() == QEvent.MouseButtonPress:
            if event.button() == Qt.LeftButton:
                try:
                    x = event.position().x()
                except Exception:
                    x = event.pos().x()
                w = le.width() or self.width()
                if self._is_rtl():
                    if x < self.ARROW_AREA:
                        self.showPopup()
                        return True
                else:
                    if x > w - self.ARROW_AREA:
                        self.showPopup()
                        return True
        return super().eventFilter(obj, event)

    def paintEvent(self, e):
        super().paintEvent(e)
        p = QPainter(self)
        p.setRenderHint(QPainter.Antialiasing)
        pm = icon_chevron_down(16).pixmap(QSize(16, 16))
        y = (self.height() - 16) // 2
        x = 4 if self._is_rtl() else self.width() - 16 - 4
        p.drawPixmap(x, y, pm)
        p.end()


# ═════════════════════════════════════════════════════════════════
#  ThemedDialog
# ═════════════════════════════════════════════════════════════════

class ThemedDialog(QDialog):
    def __init__(self, parent, title, message, kind="info",
                 buttons=None):
        if buttons is None:
            buttons = ((tr(T.ok), True, QDialog.Accepted),)
        super().__init__(parent)
        self.setWindowFlags(Qt.Dialog | Qt.FramelessWindowHint)
        self.setAttribute(Qt.WA_TranslucentBackground)
        self.setModal(True)
        self.setLayoutDirection(_dir())
        self.setFixedWidth(430)

        outer = QVBoxLayout(self)
        outer.setContentsMargins(16, 16, 16, 16)

        frame = QFrame(); frame.setObjectName("dlgFrame")
        frame.setStyleSheet(f"""
            QFrame#dlgFrame {{
                background-color: #FFFFFF;
                border-radius: 14px;
                border: 1px solid {C_BORDER};
            }}
        """)
        sh = QGraphicsDropShadowEffect()
        sh.setBlurRadius(34); sh.setOffset(0, 8)
        sh.setColor(QColor(0, 0, 0, 85))
        frame.setGraphicsEffect(sh)
        outer.addWidget(frame)

        v = QVBoxLayout(frame); v.setContentsMargins(0, 0, 0, 0); v.setSpacing(0)

        tb = QFrame(); tb.setFixedHeight(42)
        tb.setStyleSheet(f"""
            background-color: {C_DARK_GREEN};
            border-top-left-radius: 14px; border-top-right-radius: 14px;
        """)
        tb_lay = QHBoxLayout(tb)
        tb_lay.setContentsMargins(14, 0, 8, 0); tb_lay.setSpacing(8)
        t = QLabel(title)
        tf = QFont(); tf.setPointSize(10); tf.setBold(True)
        t.setFont(tf); t.setStyleSheet("color: white; background: transparent;")
        tb_lay.addWidget(t); tb_lay.addStretch()
        close_btn = QPushButton()
        close_btn.setIcon(icon_close(14, "#FFFFFF"))
        close_btn.setIconSize(QSize(14, 14))
        close_btn.setFixedSize(28, 24)
        close_btn.setCursor(Qt.PointingHandCursor)
        close_btn.setStyleSheet(f"""
            QPushButton {{
                color: white; background: transparent; border: none;
                font-size: 12px; border-radius: 5px; padding-bottom: 2px;
            }}
            QPushButton:hover {{ background-color: {C_RED}; }}
        """)
        close_btn.clicked.connect(self.reject)
        tb_lay.addWidget(close_btn)
        v.addWidget(tb)

        body = QWidget()
        body_lay = QVBoxLayout(body)
        body_lay.setContentsMargins(24, 22, 24, 18); body_lay.setSpacing(18)

        row = QHBoxLayout(); row.setSpacing(14)
        ic = QLabel()
        ic.setPixmap(icon_status(42, kind).pixmap(QSize(42, 42)))
        ic.setFixedSize(42, 42)
        ic.setStyleSheet("background: transparent;")
        ic.setAlignment(Qt.AlignTop | Qt.AlignHCenter)
        msg = QLabel(message)
        mf = QFont(); mf.setPointSize(10)
        msg.setFont(mf); msg.setWordWrap(True)
        msg.setStyleSheet(f"color: {C_DARK_GREEN}; background: transparent;")
        row.addWidget(ic); row.addWidget(msg, 1)
        body_lay.addLayout(row)

        btn_row = QHBoxLayout(); btn_row.setSpacing(10)
        btn_row.addStretch()
        for text, primary, result in buttons:
            b = QPushButton(text)
            b.setCursor(Qt.PointingHandCursor)
            b.setFixedHeight(38); b.setMinimumWidth(96)
            if primary:
                b.setStyleSheet(f"""
                    QPushButton {{
                        background-color: {C_GREEN};
                        color: white; border: none;
                        border-radius: 8px; font-size: 11px;
                        font-weight: bold; padding: 0 18px;
                    }}
                    QPushButton:hover {{ background-color: {C_DARK_GREEN}; }}
                """)
            else:
                b.setStyleSheet(f"""
                    QPushButton {{
                        background-color: transparent;
                        color: {C_DARK_GREEN};
                        border: 1px solid {C_BORDER};
                        border-radius: 8px; font-size: 11px;
                        padding: 0 18px;
                    }}
                    QPushButton:hover {{
                        background-color: {C_LIGHT_GREEN};
                        border-color: {C_GREEN};
                    }}
                """)
            b.clicked.connect(lambda _=None, r=result: self.done(r))
            btn_row.addWidget(b)
        body_lay.addLayout(btn_row)
        v.addWidget(body)

    @staticmethod
    def info(parent, title, message):
        ThemedDialog(parent, title, message, "info").exec()

    @staticmethod
    def warning(parent, title, message):
        ThemedDialog(parent, title, message, "warning").exec()

    @staticmethod
    def error(parent, title, message):
        ThemedDialog(parent, title, message, "error").exec()

    @staticmethod
    def confirm(parent, title, message):
        d = ThemedDialog(parent, title, message, "question",
                         buttons=((tr(T.yes), True, QDialog.Accepted),
                                  (tr(T.cancel), False, QDialog.Rejected)))
        return d.exec() == QDialog.Accepted

    @staticmethod
    def confirm_delete(parent, title, message):
        d = ThemedDialog(parent, title, message, "delete",
                         buttons=((tr(T.delete), True, QDialog.Accepted),
                                  (tr(T.cancel), False, QDialog.Rejected)))
        return d.exec() == QDialog.Accepted

    def showEvent(self, e):
        super().showEvent(e)
        par = self.parent()
        if par:
            pg = par.window().geometry()
            self.move(pg.center().x() - self.width() // 2,
                      pg.center().y() - self.height() // 2)


# ═════════════════════════════════════════════════════════════════
#  QuickAccessPicker
# ═════════════════════════════════════════════════════════════════

class QuickAccessPicker(QDialog):
    def __init__(self, parent, all_templates, current):
        super().__init__(parent)
        self.setWindowFlags(Qt.Dialog | Qt.FramelessWindowHint)
        self.setAttribute(Qt.WA_TranslucentBackground)
        self.setModal(True)
        self.setLayoutDirection(_dir())
        self.setFixedWidth(440)

        self._checks = {}

        outer = QVBoxLayout(self)
        outer.setContentsMargins(16, 16, 16, 16)

        frame = QFrame(); frame.setObjectName("dlgFrame")
        frame.setStyleSheet(f"""
            QFrame#dlgFrame {{
                background-color: #FFFFFF;
                border-radius: 14px;
                border: 1px solid {C_BORDER};
            }}
        """)
        sh = QGraphicsDropShadowEffect()
        sh.setBlurRadius(34); sh.setOffset(0, 8)
        sh.setColor(QColor(0, 0, 0, 85))
        frame.setGraphicsEffect(sh)
        outer.addWidget(frame)

        v = QVBoxLayout(frame); v.setContentsMargins(0, 0, 0, 0); v.setSpacing(0)

        tb = QFrame(); tb.setFixedHeight(42)
        tb.setStyleSheet(f"""
            background-color: {C_DARK_GREEN};
            border-top-left-radius: 14px; border-top-right-radius: 14px;
        """)
        tb_lay = QHBoxLayout(tb)
        tb_lay.setContentsMargins(14, 0, 8, 0); tb_lay.setSpacing(8)
        t = QLabel(tr(T.add_to_qa))
        tf = QFont(); tf.setPointSize(10); tf.setBold(True)
        t.setFont(tf); t.setStyleSheet("color: white; background: transparent;")
        tb_lay.addWidget(t); tb_lay.addStretch()
        close_btn = QPushButton()
        close_btn.setIcon(icon_close(14, "#FFFFFF"))
        close_btn.setIconSize(QSize(14, 14))
        close_btn.setFixedSize(28, 24)
        close_btn.setCursor(Qt.PointingHandCursor)
        close_btn.setStyleSheet(f"""
            QPushButton {{
                color: white; background: transparent; border: none;
                font-size: 12px; border-radius: 5px; padding-bottom: 2px;
            }}
            QPushButton:hover {{ background-color: {C_RED}; }}
        """)
        close_btn.clicked.connect(self.reject)
        tb_lay.addWidget(close_btn)
        v.addWidget(tb)

        body = QWidget()
        body_lay = QVBoxLayout(body)
        body_lay.setContentsMargins(20, 16, 20, 16); body_lay.setSpacing(12)

        hint = QLabel(tr(T.select_qa_hint))
        hint.setStyleSheet(f"color: {C_MUTED}; background: transparent; font-size: 10px;")
        hint.setWordWrap(True)
        body_lay.addWidget(hint)

        if not all_templates:
            empty = QLabel(tr(T.no_templates_yet))
            empty.setStyleSheet(
                f"color: {C_MUTED}; background: transparent; "
                f"font-size: 11px; padding: 20px;")
            empty.setAlignment(Qt.AlignCenter)
            body_lay.addWidget(empty)
        else:
            scroll = QScrollArea()
            scroll.setWidgetResizable(True)
            scroll.setFrameShape(QFrame.NoFrame)
            scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
            scroll.setStyleSheet(f"""
                QScrollArea {{ background: transparent; border: none; }}
                QScrollBar:vertical {{
                    background: transparent; width: 6px;
                }}
                QScrollBar::handle:vertical {{
                    background: {C_BORDER}; border-radius: 3px; min-height: 24px;
                }}
                QScrollBar::handle:vertical:hover {{ background: {C_GREEN}; }}
                QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical {{
                    height: 0;
                }}
            """)

            content = QWidget()
            content.setStyleSheet("background: transparent;")
            cl = QVBoxLayout(content)
            cl.setContentsMargins(0, 0, 0, 0); cl.setSpacing(4)

            for tpl in all_templates:
                name = (tpl.get("name") or "").strip()
                if not name:
                    continue
                cb = QCheckBox(name)
                cb.setChecked(name in current)
                cb.setCursor(Qt.PointingHandCursor)
                cb.setStyleSheet(f"""
                    QCheckBox {{
                        color: {C_DARK_GREEN};
                        font-size: 11px;
                        padding: 6px 8px;
                        border-radius: 6px;
                    }}
                    QCheckBox:hover {{
                        background-color: {C_HOVER};
                    }}
                    QCheckBox::indicator {{
                        width: 16px; height: 16px;
                        border: 1.5px solid {C_BORDER};
                        border-radius: 4px;
                        background: #FFFFFF;
                    }}
                    QCheckBox::indicator:hover {{
                        border-color: {C_GREEN};
                    }}
                    QCheckBox::indicator:checked {{
                        background-color: {C_GREEN};
                        border-color: {C_GREEN};
                    }}
                """)
                cl.addWidget(cb)
                self._checks[name] = cb
            cl.addStretch()

            scroll.setWidget(content)
            scroll.setMinimumHeight(200)
            scroll.setMaximumHeight(320)
            body_lay.addWidget(scroll)

        btn_row = QHBoxLayout(); btn_row.setSpacing(10)
        btn_row.addStretch()

        ok_btn = QPushButton(tr(T.save))
        ok_btn.setCursor(Qt.PointingHandCursor)
        ok_btn.setFixedHeight(38); ok_btn.setMinimumWidth(100)
        ok_btn.setStyleSheet(f"""
            QPushButton {{
                background-color: {C_GREEN}; color: white;
                border: none; border-radius: 8px;
                font-size: 11px; font-weight: bold;
                padding: 0 18px;
            }}
            QPushButton:hover {{ background-color: {C_DARK_GREEN}; }}
        """)
        ok_btn.clicked.connect(self.accept)
        btn_row.addWidget(ok_btn)

        cancel_btn = QPushButton(tr(T.cancel))
        cancel_btn.setCursor(Qt.PointingHandCursor)
        cancel_btn.setFixedHeight(38); cancel_btn.setMinimumWidth(100)
        cancel_btn.setStyleSheet(f"""
            QPushButton {{
                background-color: transparent;
                color: {C_DARK_GREEN};
                border: 1px solid {C_BORDER};
                border-radius: 8px; font-size: 11px;
                padding: 0 18px;
            }}
            QPushButton:hover {{
                background-color: {C_LIGHT_GREEN};
                border-color: {C_GREEN};
            }}
        """)
        cancel_btn.clicked.connect(self.reject)
        btn_row.addWidget(cancel_btn)
        body_lay.addLayout(btn_row)

        v.addWidget(body)

    def selected_names(self):
        return [name for name, cb in self._checks.items() if cb.isChecked()]

    def showEvent(self, e):
        super().showEvent(e)
        par = self.parent()
        if par:
            pg = par.window().geometry()
            self.move(pg.center().x() - self.width() // 2,
                      pg.center().y() - self.height() // 2)


# ═════════════════════════════════════════════════════════════════
#  JSON management
# ═════════════════════════════════════════════════════════════════

def load_templates_raw():
    if not TEMPLATES_JSON.exists():
        return []
    try:
        with open(TEMPLATES_JSON, "r", encoding="utf-8") as f:
            data = load(f)
        return data if isinstance(data, list) else []
    except Exception as e:
        print(f"[error] reading JSON: {e}")
        return []


def load_templates():
    return [t for t in load_templates_raw()
            if isinstance(t, dict) and "name" in t and "address" in t]


def save_templates(templates):
    WORDCLONER_DIR.mkdir(parents=True, exist_ok=True)
    with open(TEMPLATES_JSON, "w", encoding="utf-8") as f:
        dump(templates, f, ensure_ascii=False, indent=4)


# ═════════════════════════════════════════════════════════════════
#  Word processing
# ═════════════════════════════════════════════════════════════════

def normalize_quotes(text):
    return (text
            .replace("\u201c", '"').replace("\u201d", '"')
            .replace("\u201e", '"').replace("\u201f", '"')
            .replace("\u2033", '"'))


def _get_wrapper():
    return SETTINGS.get("placeholder_wrapper", '##')


class WrapperCountError(ValueError):
    pass


def _run_size_pt(run):
    try:
        if run.font.size is not None:
            return run.font.size.pt
    except Exception:
        pass
    return None


def compute_auto_font(runs, char_map, text, ph_start, ph_end):
    def collect_before():
        sizes = []
        i = ph_start - 1
        while i >= 0:
            while i >= 0 and text[i].isspace():
                i -= 1
            if i < 0:
                break
            while i >= 0 and not text[i].isspace():
                i -= 1
            word_start = i + 1
            ri, _ = char_map[word_start]
            s = _run_size_pt(runs[ri])
            if s is not None:
                sizes.append(s)
        return sizes

    def collect_after():
        sizes = []
        i = ph_end
        n = len(text)
        while i < n:
            while i < n and text[i].isspace():
                i += 1
            if i >= n:
                break
            word_start = i
            while i < n and not text[i].isspace():
                i += 1
            ri, _ = char_map[word_start]
            s = _run_size_pt(runs[ri])
            if s is not None:
                sizes.append(s)
        return sizes

    before = collect_before()
    after = collect_after()

    checked = []
    max_d = max(len(before), len(after))
    for d in range(max_d):
        b = before[d] if d < len(before) else None
        a = after[d] if d < len(after) else None
        if b is not None:
            checked.append(b)
        if a is not None:
            checked.append(a)
        if b is not None and a is not None and abs(b - a) < 0.01:
            return b
        if b is None or a is None:
            break

    if not checked:
        return None
    c = Counter(round(x, 1) for x in checked)
    return c.most_common(1)[0][0]


def _nearest_run(runs, char_map, text, ph_start, ph_end):
    i = ph_start - 1
    while i >= 0 and text[i].isspace():
        i -= 1
    if i >= 0:
        while i >= 0 and not text[i].isspace():
            i -= 1
        word_start = i + 1
        return runs[char_map[word_start][0]]
    i = ph_end
    n = len(text)
    while i < n and text[i].isspace():
        i += 1
    if i < n:
        return runs[char_map[i][0]]
    return None


def _copy_run_fonts(src_run, dst_run):
    if not HAS_DOCX:
        return
    try:
        src_rPr = src_run._element.find(qn('w:rPr'))
        if src_rPr is None:
            return
        src_fonts = src_rPr.find(qn('w:rFonts'))
        if src_fonts is None:
            return
        dst_rPr = dst_run._element.get_or_add_rPr()
        existing = dst_rPr.find(qn('w:rFonts'))
        if existing is not None:
            dst_rPr.remove(existing)
        dst_rPr.insert(0, deepcopy(src_fonts))
    except Exception:
        pass


def _template_paragraphs_in_document_order(doc):
    """Return body paragraphs and table-cell paragraphs in reading order.

    Placeholder replacement already processes body paragraphs and top-level
    tables. Keeping extraction in the same document order ensures CLI values
    map predictably even when a template uses table-based layouts.
    """
    paragraphs = []
    seen = set()

    def add_paragraph(paragraph):
        element = paragraph._p
        if element in seen:
            return
        seen.add(element)
        paragraphs.append(paragraph)

    for block in doc.element.body.iterchildren():
        if block.tag == qn("w:p"):
            add_paragraph(DocxParagraph(block, doc))
        elif block.tag == qn("w:tbl"):
            table = DocxTable(block, doc)
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        add_paragraph(paragraph)

    return paragraphs


def extract_placeholders(docx_path):
    doc = Document(docx_path)
    paragraphs = _template_paragraphs_in_document_order(doc)
    para_texts = [p.text for p in paragraphs]

    para_starts = []
    offset = 0
    for t in para_texts:
        para_starts.append(offset)
        offset += len(t) + 1

    raw_text = "\n".join(para_texts)
    text = normalize_quotes(raw_text)

    wrapper = normalize_quotes(_get_wrapper())
    wlen = len(wrapper)
    if wlen == 0:
        return {}

    positions = []
    start = 0
    while True:
        idx = text.find(wrapper, start)
        if idx == -1:
            break
        positions.append(idx)
        start = idx + wlen

    if len(positions) % 2 != 0:
        raise WrapperCountError(tr(T.var_count_error_detail,
                                    w=wrapper, n=len(positions)))

    props = {}
    for i in range(0, len(positions) - 1, 2):
        open_pos = positions[i]
        close_pos = positions[i + 1]
        term_start = open_pos + wlen
        term = text[term_start:close_pos]
        term_key = (term or "").strip()
        if not term_key or term_key in props:
            continue

        pi = 0
        for k in range(len(para_starts) - 1, -1, -1):
            if para_starts[k] <= open_pos:
                pi = k
                break

        auto_size = None
        font_name = None

        if pi < len(para_texts):
            para_start = para_starts[pi]
            para_text_norm = normalize_quotes(para_texts[pi])
            local_start = open_pos - para_start
            local_end = close_pos + wlen - para_start

            if (0 <= local_start < len(para_text_norm)
                    and local_end <= len(para_text_norm)):
                p = paragraphs[pi]
                runs = list(p.runs)
                if runs:
                    parts = []
                    char_map = []
                    for ri, r in enumerate(runs):
                        for ci, ch in enumerate(r.text):
                            parts.append(ch)
                            char_map.append((ri, ci))
                    para_join = normalize_quotes("".join(parts))
                    if len(para_join) == len(para_text_norm):
                        try:
                            auto_size = compute_auto_font(
                                runs, char_map, para_join,
                                local_start, local_end)
                        except Exception:
                            auto_size = None
                        try:
                            ref = _nearest_run(
                                runs, char_map, para_join,
                                local_start, local_end)
                            if ref is not None:
                                try:
                                    if ref.font.name:
                                        font_name = ref.font.name
                                except Exception:
                                    pass
                        except Exception:
                            pass

        props[term_key] = {
            "start": term_start,
            "auto_size": auto_size,
            "font_name": font_name,
        }
    return props


def replace_placeholders_in_docx(src_path, dst_path, replacements, font_overrides):
    doc = Document(src_path)
    wrapper = normalize_quotes(_get_wrapper())
    wlen = len(wrapper)
    if wlen == 0:
        doc.save(dst_path)
        return

    def process_paragraph(paragraph):
        original_runs = list(paragraph.runs)
        if not original_runs:
            return

        parts = []
        char_map = []
        for ri, r in enumerate(original_runs):
            for ci, ch in enumerate(r.text):
                parts.append(ch)
                char_map.append((ri, ci))
        text = "".join(parts)
        norm = normalize_quotes(text)

        ranges = []
        i = 0
        while i < len(norm):
            start = norm.find(wrapper, i)
            if start == -1:
                break
            end = norm.find(wrapper, start + wlen)
            if end == -1:
                break
            key = norm[start + wlen:end].strip()
            if key in replacements:
                ranges.append((start, end + wlen, key))
                i = end + wlen
            else:
                i = start + wlen

        if not ranges:
            return

        auto_sizes = {}
        auto_refs = {}
        for (start, end, key) in ranges:
            override = font_overrides.get(key, (None, False, False, None))
            if override[0] is None:
                auto_sizes[(start, end, key)] = compute_auto_font(
                    original_runs, char_map, norm, start, end)
            else:
                auto_sizes[(start, end, key)] = override[0]
            auto_refs[(start, end, key)] = _nearest_run(
                original_runs, char_map, norm, start, end)

        segments = []

        def emit_orig(lo, hi):
            if lo >= hi:
                return
            seg_lo = lo
            cur_ri = char_map[lo][0]
            for k in range(lo + 1, hi):
                if char_map[k][0] != cur_ri:
                    segments.append((norm[seg_lo:k], original_runs[cur_ri], False))
                    seg_lo = k
                    cur_ri = char_map[k][0]
            segments.append((norm[seg_lo:hi], original_runs[cur_ri], False))

        pos = 0
        for (start, end, key) in ranges:
            emit_orig(pos, start)
            override = font_overrides.get(key, (None, False, False, None))
            size_pt = auto_sizes[(start, end, key)]
            ref_run = auto_refs.get((start, end, key))
            color_hex = override[3] if len(override) > 3 else None
            value = replacements[key]
            segments.append(
                (value,
                 (size_pt, override[1], override[2], color_hex, ref_run),
                 True)
            )
            pos = end

        if pos < len(norm):
            emit_orig(pos, len(norm))

        for r in original_runs:
            r._element.getparent().remove(r._element)

        for seg_text, font_spec, is_repl in segments:
            if not seg_text:
                continue
            new_run = paragraph.add_run(seg_text)
            if is_repl:
                size_pt, bold, italic, color_hex, ref_run = font_spec
                if ref_run is not None:
                    _copy_run_fonts(ref_run, new_run)
                if size_pt is not None:
                    try:
                        new_run.font.size = Pt(float(size_pt))
                    except Exception:
                        pass
                try:
                    new_run.font.bold = bool(bold)
                    new_run.font.italic = bool(italic)
                except Exception:
                    pass
                if color_hex:
                    try:
                        h = color_hex.lstrip("#").upper()
                        new_run.font.color.rgb = RGBColor.from_string(h)
                    except Exception:
                        pass
            else:
                orig = font_spec
                try:
                    new_run.bold = orig.bold
                    new_run.italic = orig.italic
                    new_run.underline = orig.underline
                    if orig.font.size is not None:
                        new_run.font.size = orig.font.size
                    if orig.font.name is not None:
                        new_run.font.name = orig.font.name
                    if (orig.font.color is not None
                            and orig.font.color.rgb is not None):
                        new_run.font.color.rgb = orig.font.color.rgb
                except Exception:
                    pass

    for p in doc.paragraphs:
        process_paragraph(p)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    process_paragraph(p)

    doc.save(dst_path)


class ExcelImportError(ValueError):
    """A workbook cannot be mapped safely to the current template."""


def _excel_is_blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _excel_cell_text(value):
    """Convert Excel cell values to stable, human-readable replacement text."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        if not any((value.hour, value.minute, value.second, value.microsecond)):
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def read_excel_records(excel_path, variable_names, header_row=False):
    """Read the active worksheet using filename + ordered-variable columns."""
    if not HAS_EXCEL or load_workbook is None:
        raise ExcelImportError(tr(T.batch_excel_library_missing))

    path = Path(excel_path).expanduser()
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        raise ExcelImportError(tr(T.batch_excel_type_error, path=path))
    if not path.is_file():
        raise ExcelImportError(tr(T.batch_excel_not_found, path=path))

    try:
        workbook = load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise ExcelImportError(tr(T.batch_excel_read_error, err=exc)) from exc

    records = []
    try:
        worksheet = workbook.active
        first_data_row = 2 if header_row else 1
        expected_columns = len(variable_names) + 1
        max_column = max(int(worksheet.max_column or 1), expected_columns)

        for row_number, row in enumerate(
                worksheet.iter_rows(
                    min_row=first_data_row,
                    max_row=int(worksheet.max_row or 0),
                    max_col=max_column,
                    values_only=True,
                ),
                start=first_data_row):
            values = list(row)
            if all(_excel_is_blank(value) for value in values):
                continue

            for offset, extra_value in enumerate(values[expected_columns:]):
                if not _excel_is_blank(extra_value):
                    extra_column = expected_columns + offset + 1
                    raise ExcelImportError(tr(
                        T.batch_extra_columns,
                        row=row_number,
                        column=extra_column,
                        variables=len(variable_names),
                    ))

            output_name = _excel_cell_text(values[0]).strip()
            if not output_name:
                raise ExcelImportError(tr(T.batch_filename_empty, row=row_number))

            replacements = {
                variable: _excel_cell_text(
                    values[index + 1] if index + 1 < len(values) else None
                )
                for index, variable in enumerate(variable_names)
            }
            records.append((row_number, output_name, replacements))
    except ExcelImportError:
        raise
    except Exception as exc:
        raise ExcelImportError(tr(T.batch_excel_read_error, err=exc)) from exc
    finally:
        try:
            workbook.close()
        except Exception:
            pass

    if not records:
        raise ExcelImportError(tr(T.batch_no_data_rows))
    return records


def _save_docx_atomically(src_path, dst_path, replacements, font_overrides,
                          overwrite=False):
    """Write a complete DOCX to a sibling temporary file before committing."""
    dst_path = Path(dst_path)
    dst_path.parent.mkdir(parents=True, exist_ok=True)
    if dst_path.exists() and not overwrite:
        raise FileExistsError(str(dst_path))

    fd, temp_name = tempfile.mkstemp(
        prefix=f".{dst_path.stem}.", suffix=".docx", dir=str(dst_path.parent)
    )
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        replace_placeholders_in_docx(
            str(src_path), str(temp_path), replacements, font_overrides
        )
        if dst_path.exists() and not overwrite:
            raise FileExistsError(str(dst_path))
        os.replace(str(temp_path), str(dst_path))
    finally:
        try:
            temp_path.unlink(missing_ok=True)
        except OSError:
            pass
    return dst_path


def _safe_batch_filename(raw_name, output_dir, reserved_names):
    name = str(raw_name).strip()
    invalid = '<>:"/\\|?*'
    name = "".join("_" if ch in invalid or ord(ch) < 32 else ch for ch in name)
    name = name.rstrip(" .")
    if not name:
        raise ValueError(tr(T.batch_filename_empty, row="?"))
    if not name.lower().endswith(".docx"):
        name += ".docx"

    candidate = Path(name)
    stem = candidate.stem.rstrip(" .") or "Document"
    suffix = ".docx"
    if stem.upper() in {"CON", "PRN", "AUX", "NUL", *(f"COM{i}" for i in range(1, 10)), *(f"LPT{i}" for i in range(1, 10))}:
        stem = f"_{stem}"

    output_dir = Path(output_dir)
    number = 1
    while True:
        file_name = f"{stem}{suffix}" if number == 1 else f"{stem}_{number}{suffix}"
        identity = file_name.casefold()
        if identity not in reserved_names and not (output_dir / file_name).exists():
            reserved_names.add(identity)
            return file_name
        number += 1


def _font_size_value(size_combo):
    size_text = size_combo.currentText().strip().lower()
    if size_text in ("", "auto", tr(T.auto).strip().lower()):
        return None
    digit_map = str.maketrans(
        "۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩",
        "01234567890123456789",
    )
    normalized = size_text.translate(digit_map).replace(",", ".").replace("٫", ".")
    try:
        return float(normalized)
    except ValueError:
        return None


# ═════════════════════════════════════════════════════════════════
#  TitleBar — buttons ALWAYS on the right, both languages
# ═════════════════════════════════════════════════════════════════

class TitleBar(QFrame):
    def __init__(self, window):
        super().__init__()
        self._win = window
        self._drag = None
        self.setFixedHeight(46)
        self.setObjectName("titleBar")

        # ── Force LTR layout so buttons stay on the right ──
        # (otherwise RTL app direction would mirror them to the left)
        self.setLayoutDirection(Qt.LeftToRight)
        
        lay = QHBoxLayout(self)
        lay.setContentsMargins(16, 0, 8, 0)
        lay.setSpacing(10)

        # ── Left: icon + title ──
        ico = QLabel()
        ico.setPixmap(icon_document(20, "#FFFFFF").pixmap(QSize(20, 20)))
        ico.setStyleSheet("background: transparent;")

        t = QLabel(tr(T.app_title))
        f = QFont(); f.setPointSize(11); f.setBold(True)
        t.setFont(f)
        t.setStyleSheet("color: white; background: transparent;")

        lay.addWidget(ico)
        lay.addWidget(t)
        lay.addStretch()

        # ── Right: minimize / maximize / close ──
        b_min = self._btn("", C_GREEN)
        b_max = self._btn("", C_GREEN)
        b_cls = self._btn("", C_RED)
        b_min.setIcon(icon_minimize(16, "#FFFFFF"))
        b_max.setIcon(icon_maximize(16, "#FFFFFF"))
        b_cls.setIcon(icon_close(14, "#FFFFFF"))
        b_min.setIconSize(QSize(16, 16))
        b_max.setIconSize(QSize(16, 16))
        b_cls.setIconSize(QSize(14, 14))
        b_min.clicked.connect(self._win.showMinimized)
        b_max.clicked.connect(self._toggle_max)
        b_cls.clicked.connect(self._win.close)

        lay.addWidget(b_min)
        lay.addWidget(b_max)
        lay.addWidget(b_cls)

        self._apply_style()

    def _apply_style(self):
        r = 0 if self._win.isMaximized() else 14
        self.setStyleSheet(f"""
            QFrame#titleBar {{
                background-color: {C_DARK_GREEN};
                border-top-left-radius: {r}px;
                border-top-right-radius: {r}px;
            }}
        """)

    def refresh_radius(self):
        self._apply_style()

    def _btn(self, txt, hover):
        b = QPushButton(txt)
        b.setFixedSize(34, 28)
        b.setCursor(Qt.PointingHandCursor)
        b.setStyleSheet(f"""
            QPushButton {{
                color: white; background: transparent; border: none;
                font-size: 14px; border-radius: 6px; padding-bottom: 2px;
            }}
            QPushButton:hover {{ background-color: {hover}; }}
        """)
        return b

    def _toggle_max(self):
        self._win.showNormal() if self._win.isMaximized() else self._win.showMaximized()

    def mousePressEvent(self, e):
        if e.button() != Qt.LeftButton:
            return
        try:
            wh = self._win.windowHandle()
            if wh is not None:
                if self._win.isMaximized():
                    self._win.showNormal()
                if wh.startSystemMove():
                    e.accept()
                    return
        except Exception:
            pass
        self._drag = e.globalPosition().toPoint() - self._win.frameGeometry().topLeft()

    def mouseMoveEvent(self, e):
        if (e.buttons() & Qt.LeftButton and self._drag
                and not self._win.isMaximized()):
            self._win.move(e.globalPosition().toPoint() - self._drag)

    def mouseReleaseEvent(self, e):
        self._drag = None

    def mouseDoubleClickEvent(self, e):
        self._toggle_max()

# ═════════════════════════════════════════════════════════════════
#  DropdownButton / DragGhost / GripLabel / DragController
# ═════════════════════════════════════════════════════════════════

class DropdownButton(QFrame):
    clicked = Signal()

    def __init__(self):
        super().__init__()
        self.setObjectName("ddbtn")
        self.setFixedHeight(58)
        self.setCursor(Qt.PointingHandCursor)
        self.setStyleSheet(f"""
            QFrame#ddbtn {{
                background-color: #FFFFFF;
                border: 1.5px solid {C_BORDER};
                border-radius: 12px;
            }}
            QFrame#ddbtn:hover {{
                border-color: {C_GREEN};
                background-color: #FBFDFC;
            }}
        """)
        lay = QHBoxLayout(self)
        lay.setContentsMargins(18, 0, 18, 0); lay.setSpacing(12)

        ico = QLabel(); ico.setPixmap(icon_document(22).pixmap(QSize(22, 22)))
        ico.setStyleSheet("background: transparent;")
        ico.setAttribute(Qt.WA_TransparentForMouseEvents)

        self.text = QLabel(tr(T.select_template))
        f = QFont(); f.setPointSize(11)
        self.text.setFont(f)
        self.text.setStyleSheet(f"color: {C_DARK_GREEN}; background: transparent;")
        self.text.setAttribute(Qt.WA_TransparentForMouseEvents)

        chev = QLabel(); chev.setPixmap(icon_chevron_down(16).pixmap(QSize(16, 16)))
        chev.setStyleSheet("background: transparent;")
        chev.setAttribute(Qt.WA_TransparentForMouseEvents)

        lay.addWidget(ico); lay.addWidget(self.text)
        lay.addStretch(); lay.addWidget(chev)

    def set_text(self, s):
        self.text.setText(s)

    def mousePressEvent(self, e):
        if e.button() == Qt.LeftButton:
            self.clicked.emit()


class DragGhost(QWidget):
    def __init__(self, text, icon, width):
        super().__init__(
            None,
            Qt.Tool | Qt.FramelessWindowHint | Qt.WindowStaysOnTopHint
        )
        self.setAttribute(Qt.WA_TranslucentBackground, True)
        self.setAttribute(Qt.WA_ShowWithoutActivating, True)
        self.setAttribute(Qt.WA_TransparentForMouseEvents, True)

        self.setFixedSize(width, 46)

        frame = QFrame(self)
        frame.setGeometry(0, 0, width, 46)
        frame.setStyleSheet(f"""
            QFrame {{
                background-color: #FFFFFF;
                border: 1.5px solid {C_GREEN};
                border-radius: 8px;
            }}
        """)
        lay = QHBoxLayout(frame)
        lay.setContentsMargins(10, 0, 10, 0); lay.setSpacing(8)

        ic = QLabel()
        ic.setPixmap(icon.pixmap(QSize(20, 20)))
        ic.setStyleSheet("background: transparent; border: none;")
        lay.addWidget(ic)

        tx = QLabel(text)
        tx.setStyleSheet(
            f"color: {C_DARK_GREEN}; background: transparent; "
            f"border: none; font-size: 11px;")
        lay.addWidget(tx)
        lay.addStretch()

        self.setWindowOpacity(0.95)


class GripLabel(QLabel):
    def __init__(self, controller, item, parent=None):
        super().__init__(parent)
        self._controller = controller
        self._item = item
        self._dragging = False
        self.setPixmap(icon_grip(16).pixmap(QSize(16, 16)))
        self.setFixedSize(18, 18)
        self.setStyleSheet("background: transparent;")
        self.setCursor(Qt.OpenHandCursor)
        self.setToolTip(tr(T.drag_to_reorder))

    def mousePressEvent(self, e):
        if e.button() == Qt.LeftButton and self._controller:
            self._dragging = True
            self.setCursor(Qt.ClosedHandCursor)
            self._controller.start(self._item, e.globalPosition().toPoint())

    def mouseMoveEvent(self, e):
        if self._dragging and self._controller:
            self._controller.move(e.globalPosition().toPoint())

    def mouseReleaseEvent(self, e):
        if self._dragging and self._controller:
            self._dragging = False
            self.setCursor(Qt.OpenHandCursor)
            self._controller.end()


class DragController:
    ITEM_H = 48

    def __init__(self, on_reorder=None):
        self.items = []
        self.layout = None
        self.on_reorder = on_reorder
        self.drag_item = None
        self.ghost = None
        self.indicator = None
        self._target_idx = None

    def setup(self, items, layout):
        self.items = items
        self.layout = layout

    def _item_name(self, item):
        try:
            tpl = getattr(item, "_tpl", None) or {}
            return tpl.get("name", "") or getattr(item, "_name", "")
        except Exception:
            return ""

    def start(self, item, global_pos):
        self.drag_item = item
        content = item.parentWidget()
        if content is None:
            return

        name = self._item_name(item)
        self.ghost = DragGhost(name, icon_document(20), item.width())
        self.ghost.move(global_pos.x() - self.ghost.width() // 2,
                        global_pos.y() - 23)
        self.ghost.show()
        self.ghost.raise_()

        self.indicator = QFrame(content)
        self.indicator.setStyleSheet(
            f"background-color: {C_GREEN}; border: none; border-radius: 1px;")
        self.indicator.setFixedHeight(2)
        self.indicator.hide()

        item.setStyleSheet(f"""
            QFrame#it {{
                background-color: {C_HOVER};
                border-radius: 8px;
                border: 1px dashed {C_BORDER};
            }}
            QFrame#qaitem {{
                background-color: {C_HOVER};
                border-radius: 10px;
                border: 1px dashed {C_BORDER};
            }}
        """)

    def move(self, global_pos):
        if not self.drag_item or self.ghost is None:
            return

        self.ghost.move(global_pos.x() - self.ghost.width() // 2,
                        global_pos.y() - 23)

        content = self.drag_item.parentWidget()
        if content is None:
            return
        local = content.mapFromGlobal(global_pos)
        mouse_y = local.y()
        n = len(self.items)
        target_idx = max(0, min(n, int((mouse_y + self.ITEM_H // 2) // self.ITEM_H)))
        self._target_idx = target_idx

        y = max(0, target_idx * self.ITEM_H - 1)
        self.indicator.setGeometry(4, y, max(10, content.width() - 8), 2)
        self.indicator.show()
        self.indicator.raise_()

    def end(self):
        if not self.drag_item:
            return

        target_idx = self._target_idx

        if self.ghost is not None:
            self.ghost.hide()
            self.ghost.deleteLater()
            self.ghost = None
        if self.indicator is not None:
            self.indicator.hide()
            self.indicator.deleteLater()
            self.indicator = None

        self.drag_item.setStyleSheet(f"""
            QFrame#it {{ background-color: transparent; border-radius: 8px; }}
            QFrame#it:hover {{ background-color: {C_HOVER}; }}
            QFrame#qaitem {{
                background-color: #FFFFFF;
                border: 1.5px solid {C_BORDER};
                border-radius: 10px;
            }}
            QFrame#qaitem:hover {{
                border-color: {C_GREEN};
                background-color: {C_HOVER};
            }}
        """)

        if target_idx is not None:
            self._move_item_to(self.drag_item, target_idx)
            if self.on_reorder:
                self.on_reorder(list(self.items))

        self.drag_item = None
        self._target_idx = None

    def _move_item_to(self, item, target_idx):
        items = self.items
        try:
            cur_idx = items.index(item)
        except ValueError:
            return
        if target_idx > cur_idx:
            target_idx -= 1
        if target_idx == cur_idx:
            return
        items.pop(cur_idx)
        items.insert(target_idx, item)
        for it in items:
            self.layout.removeWidget(it)
        for i, it in enumerate(items):
            self.layout.insertWidget(i, it)


class QuickAccessDragController(DragController):
    ITEM_H = 44


class PopupItem(QFrame):
    clicked = Signal()
    delete_clicked = Signal()

    def __init__(self, icon: QIcon, text: str, *, is_add=False,
                 show_delete=False, show_grip=False,
                 drag_controller=None):
        super().__init__()
        self.setObjectName("it")
        self.setFixedHeight(46)
        self.setCursor(Qt.PointingHandCursor)
        self._tpl = None
        self.grip_label = None

        bg       = C_LIGHT_GREEN if is_add else "transparent"
        bg_hover = C_ACCENT1 if is_add else C_HOVER
        self.setStyleSheet(f"""
            QFrame#it {{ background-color: {bg}; border-radius: 8px; }}
            QFrame#it:hover {{ background-color: {bg_hover}; }}
        """)
        lay = QHBoxLayout(self)
        lay.setContentsMargins(8, 0, 10, 0); lay.setSpacing(8)

        if show_grip and drag_controller is not None:
            grip = GripLabel(drag_controller, self)
            self.grip_label = grip
            lay.addWidget(grip)

        ic = QLabel(); ic.setPixmap(icon.pixmap(QSize(26, 26)))
        ic.setStyleSheet("background: transparent;")
        ic.setAttribute(Qt.WA_TransparentForMouseEvents)
        lay.addWidget(ic)

        tx = QLabel(text)
        f = QFont(); f.setPointSize(11)
        if is_add: f.setBold(True)
        tx.setFont(f)
        tx.setStyleSheet(f"color: {C_DARK_GREEN}; background: transparent;")
        tx.setAttribute(Qt.WA_TransparentForMouseEvents)
        lay.addWidget(tx)
        lay.addStretch()

        self.del_btn = None
        if show_delete:
            self.del_btn = QPushButton()
            self.del_btn.setIcon(icon_trash(18))
            self.del_btn.setIconSize(QSize(18, 18))
            self.del_btn.setFixedSize(30, 30)
            self.del_btn.setCursor(Qt.PointingHandCursor)
            self.del_btn.setVisible(False)
            self.del_btn.setStyleSheet("""
                QPushButton { background: transparent; border: none; border-radius: 6px; }
                QPushButton:hover { background-color: #FCE8E6; }
            """)
            self.del_btn.clicked.connect(self.delete_clicked.emit)
            lay.addWidget(self.del_btn)

    def enterEvent(self, e):
        if self.del_btn: self.del_btn.setVisible(True)

    def leaveEvent(self, e):
        if self.del_btn: self.del_btn.setVisible(False)

    def mousePressEvent(self, e):
        if e.button() != Qt.LeftButton: return
        if self.del_btn and self.del_btn.isVisible():
            if self.del_btn.geometry().contains(e.position().toPoint()):
                return
        self.clicked.emit()


class DropdownPopup(QWidget):
    add_clicked = Signal()
    template_selected = Signal(dict)
    template_deleted = Signal(dict)
    reordered = Signal(list)

    ADD_ITEM_H = 46
    SEP_H      = 10
    PADDING_H  = 16
    SEARCH_H   = 44
    TEMPLATE_H = 48

    def __init__(self, templates, parent=None):
        super().__init__(
            parent,
            Qt.Popup | Qt.FramelessWindowHint | Qt.NoDropShadowWindowHint
        )
        self.setAttribute(Qt.WA_TranslucentBackground)
        self.setLayoutDirection(_dir())

        self._templates = templates
        self._scroll = None
        self._content_h = 0
        self._items = []
        self._max_scroll_h = 320

        self._drag = DragController(on_reorder=self._on_reorder)

        outer = QVBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)

        frame = QFrame(); frame.setObjectName("popupFrame")
        frame.setStyleSheet(f"""
            QFrame#popupFrame {{
                background-color: #FFFFFF;
                border: 1px solid {C_BORDER};
                border-radius: 14px;
            }}
        """)
        outer.addWidget(frame)

        inner = QVBoxLayout(frame)
        inner.setContentsMargins(8, 8, 8, 8); inner.setSpacing(2)

        add_item = PopupItem(icon_plus_circle(28), tr(T.add_letter), is_add=True)
        add_item.clicked.connect(self.add_clicked.emit)
        inner.addWidget(add_item)

        search_input = QLineEdit()
        search_input.setPlaceholderText(tr(T.search_template))
        search_input.setFixedHeight(34)
        search_input.setLayoutDirection(_dir())
        search_input.setStyleSheet(f"""
            QLineEdit {{
                background-color: {C_HOVER};
                border: 1px solid {C_BORDER};
                border-radius: 8px;
                padding: 0 10px;
                color: {C_DARK_GREEN};
                font-size: 11px;
                margin: 4px 2px 4px 2px;
            }}
            QLineEdit:focus {{
                border-color: {C_GREEN};
                background-color: #FFFFFF;
            }}
        """)
        search_input.textChanged.connect(self._on_search_changed)
        inner.addWidget(search_input)
        self._search = search_input

        if templates:
            sep = QFrame(); sep.setFixedHeight(1)
            sep.setStyleSheet(f"background-color: {C_BORDER}; margin: 2px 10px;")
            inner.addWidget(sep)

            scroll = QScrollArea()
            scroll.setWidgetResizable(True)
            scroll.setFrameShape(QFrame.NoFrame)
            scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
            scroll.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
            scroll.setStyleSheet(f"""
                QScrollArea {{ background: transparent; border: none; }}
                QScrollArea > QWidget > QWidget {{ background: transparent; }}
                QScrollBar:vertical {{
                    background: transparent; width: 6px; margin: 2px 0 2px 0;
                    border: none;
                }}
                QScrollBar::handle:vertical {{
                    background: {C_BORDER}; border-radius: 3px; min-height: 24px;
                }}
                QScrollBar::handle:vertical:hover {{ background: {C_GREEN}; }}
                QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical {{
                    height: 0; background: transparent;
                }}
                QScrollBar::add-page:vertical, QScrollBar::sub-page:vertical {{
                    background: transparent;
                }}
            """)

            content = QWidget()
            content.setStyleSheet("background: transparent;")
            cl = QVBoxLayout(content)
            cl.setContentsMargins(0, 0, 0, 0); cl.setSpacing(2)

            for tpl in templates:
                it = PopupItem(
                    icon_document(22), tpl.get("name", tr(T.untitled)),
                    show_delete=True, show_grip=True,
                    drag_controller=self._drag)
                it._tpl = tpl
                it.clicked.connect(lambda _=None, t=tpl: self.template_selected.emit(t))
                it.delete_clicked.connect(lambda _=None, t=tpl: self.template_deleted.emit(t))
                cl.addWidget(it)
                self._items.append(it)
            cl.addStretch()

            self._drag.setup(self._items, cl)

            scroll.setWidget(content)
            self._scroll = scroll
            self._content_h = len(templates) * self.TEMPLATE_H + 8
            scroll.setFixedHeight(self._content_h)
            inner.addWidget(scroll)

    def _on_search_changed(self, text):
        text = (text or "").strip().lower()
        is_filtering = bool(text)
        count = 0
        for item in self._items:
            name = ""
            if item._tpl is not None:
                name = (item._tpl.get("name") or "").lower()
            show = (not text) or (text in name)
            item.setVisible(show)
            if show:
                count += 1
            if item.grip_label is not None:
                item.grip_label.setEnabled(not is_filtering)
                item.grip_label.setVisible(not is_filtering)

        if self._scroll is not None:
            new_h = count * self.TEMPLATE_H + 8 if count > 0 else 0
            self._content_h = new_h
            final_h = min(new_h, self._max_scroll_h)
            self._scroll.setFixedHeight(final_h)

    def _on_reorder(self, items):
        new_order = [it._tpl for it in items if it._tpl is not None]
        if not new_order:
            return
        try:
            save_templates(new_order)
        except Exception as e:
            print(f"[error] saving order: {e}")
        self.reordered.emit(new_order)

    def constrain_height(self, available_h):
        if self._scroll is None:
            return
        overhead = (self.ADD_ITEM_H + self.SEP_H + self.PADDING_H
                    + self.SEARCH_H)
        max_scroll = max(60, int(available_h) - overhead)
        self._max_scroll_h = max_scroll
        final_h = min(self._content_h, max_scroll)
        self._scroll.setFixedHeight(final_h)


class TemplateDropdown(QWidget):
    add_requested = Signal()
    template_selected = Signal(dict)
    template_deleted = Signal(dict)

    def __init__(self):
        super().__init__()
        self._templates = []; self._popup = None
        lay = QVBoxLayout(self); lay.setContentsMargins(0, 0, 0, 0)
        self.button = DropdownButton()
        self.button.clicked.connect(self._toggle)
        lay.addWidget(self.button)

    def set_templates(self, tpls):
        self._templates = tpls

    def _toggle(self):
        if self._popup and self._popup.isVisible():
            self._popup.close(); self._popup = None; return

        popup = DropdownPopup(self._templates, self)
        popup.add_clicked.connect(self._on_add)
        popup.template_selected.connect(self._on_select)
        popup.template_deleted.connect(self._on_delete)
        popup.reordered.connect(self._on_reordered)

        popup.setFixedWidth(self.button.width())

        gp = self.button.mapToGlobal(QPoint(0, self.button.height() + 4))

        screen = QGuiApplication.screenAt(gp) or QGuiApplication.primaryScreen()
        avail = screen.availableGeometry()
        available_h = avail.bottom() - gp.y() - 10
        popup.constrain_height(available_h)

        popup.adjustSize()
        popup.move(gp.x(), gp.y())
        popup.show()
        self._popup = popup

    def _close(self):
        if self._popup:
            self._popup.close(); self._popup = None

    def _on_add(self):
        self._close(); self.add_requested.emit()

    def _on_select(self, tpl):
        self._close()
        self.button.set_text(tpl.get("name", tr(T.select_template)))
        self.template_selected.emit(tpl)

    def _on_delete(self, tpl):
        self._close(); self.template_deleted.emit(tpl)

    def _on_reordered(self, new_list):
        self._templates = list(new_list)


# ═════════════════════════════════════════════════════════════════
#  Quick Access
# ═════════════════════════════════════════════════════════════════

class QuickAccessItem(QFrame):
    clicked = Signal()
    removed = Signal()

    def __init__(self, name, drag_controller=None):
        super().__init__()
        self.setObjectName("qaitem")
        self.setFixedHeight(40)
        self.setCursor(Qt.PointingHandCursor)
        self._name = name
        self._tpl = {"name": name}
        self.grip_label = None

        self.setStyleSheet(f"""
            QFrame#qaitem {{
                background-color: #FFFFFF;
                border: 1.5px solid {C_BORDER};
                border-radius: 10px;
            }}
            QFrame#qaitem:hover {{
                border-color: {C_GREEN};
                background-color: {C_HOVER};
            }}
        """)
        lay = QHBoxLayout(self)
        lay.setContentsMargins(4, 0, 8, 0); lay.setSpacing(6)

        if drag_controller is not None:
            grip = GripLabel(drag_controller, self)
            self.grip_label = grip
            lay.addWidget(grip)

        ic = QLabel()
        ic.setPixmap(icon_document(16).pixmap(QSize(16, 16)))
        ic.setStyleSheet("background: transparent;")
        ic.setAttribute(Qt.WA_TransparentForMouseEvents)
        lay.addWidget(ic)

        tx = QLabel(name)
        f = QFont(); f.setPointSize(10)
        tx.setFont(f)
        tx.setStyleSheet(f"color: {C_DARK_GREEN}; background: transparent;")
        tx.setAttribute(Qt.WA_TransparentForMouseEvents)
        lay.addWidget(tx)
        lay.addStretch()

        self.del_btn = QPushButton()
        self.del_btn.setIcon(icon_close(12, C_MUTED))
        self.del_btn.setIconSize(QSize(12, 12))
        self.del_btn.setFixedSize(22, 22)
        self.del_btn.setCursor(Qt.PointingHandCursor)
        self.del_btn.setToolTip(tr(T.remove_from_qa))
        self.del_btn.setVisible(False)
        self.del_btn.setStyleSheet(f"""
            QPushButton {{
                color: {C_MUTED};
                background: transparent;
                border: none;
                font-size: 11px;
                border-radius: 11px;
            }}
            QPushButton:hover {{
                background-color: {C_RED};
                color: white;
            }}
        """)
        self.del_btn.clicked.connect(self.removed.emit)
        lay.addWidget(self.del_btn)

    def enterEvent(self, e):
        self.del_btn.setVisible(True)

    def leaveEvent(self, e):
        self.del_btn.setVisible(False)

    def mousePressEvent(self, e):
        if e.button() != Qt.LeftButton:
            return
        if self.del_btn.isVisible() and \
                self.del_btn.geometry().contains(e.position().toPoint()):
            return
        self.clicked.emit()


class QuickAccessWidget(QWidget):
    template_selected = Signal(dict)
    add_requested = Signal()
    quick_access_changed = Signal()

    def __init__(self):
        super().__init__()
        self.setStyleSheet("background: transparent;")
        self._templates = []
        self._quick_names = []

        self._drag = QuickAccessDragController(on_reorder=self._on_reorder)

        lay = QVBoxLayout(self)
        lay.setContentsMargins(0, 0, 0, 0); lay.setSpacing(6)

        header = QHBoxLayout()
        header.setContentsMargins(0, 0, 0, 0); header.setSpacing(8)

        title = QLabel(tr(T.quick_access))
        tf = QFont(); tf.setPointSize(11); tf.setBold(True)
        title.setFont(tf)
        title.setStyleSheet(f"color: {C_DARK_GREEN}; background: transparent;")
        header.addWidget(title)
        header.addStretch()

        self.add_btn = QPushButton()
        self.add_btn.setIcon(icon_plus_white(14))
        self.add_btn.setIconSize(QSize(14, 14))
        self.add_btn.setFixedSize(26, 26)
        self.add_btn.setCursor(Qt.PointingHandCursor)
        self.add_btn.setToolTip(tr(T.add_template_to_qa))
        self.add_btn.setStyleSheet(f"""
            QPushButton {{
                background-color: {C_GREEN};
                border: none;
                border-radius: 13px;
                padding: 0px;
            }}
            QPushButton:hover {{ background-color: {C_DARK_GREEN}; }}
        """)
        self.add_btn.clicked.connect(self.add_requested.emit)
        header.addWidget(self.add_btn)
        lay.addLayout(header)

        self.scroll = QScrollArea()
        self.scroll.setWidgetResizable(True)
        self.scroll.setFrameShape(QFrame.NoFrame)
        self.scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.scroll.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        self.scroll.setStyleSheet(f"""
            QScrollArea {{ background: transparent; border: none; }}
            QScrollArea > QWidget > QWidget {{ background: transparent; }}
            QScrollBar:vertical {{
                background: transparent; width: 6px;
            }}
            QScrollBar::handle:vertical {{
                background: {C_BORDER}; border-radius: 3px; min-height: 24px;
            }}
            QScrollBar::handle:vertical:hover {{ background: {C_GREEN}; }}
            QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical {{
                height: 0;
            }}
        """)

        self.items_container = QWidget()
        self.items_container.setStyleSheet("background: transparent;")
        self.items_layout = QVBoxLayout(self.items_container)
        self.items_layout.setContentsMargins(0, 0, 4, 0)
        self.items_layout.setSpacing(4)
        self.items_layout.addStretch()

        self.scroll.setWidget(self.items_container)
        lay.addWidget(self.scroll)

        self.empty_lbl = QLabel(tr(T.no_qa_templates))
        self.empty_lbl.setAlignment(Qt.AlignCenter)
        self.empty_lbl.setStyleSheet(
            f"color: {C_MUTED}; background: transparent; "
            f"font-size: 10px; padding: 10px;")
        lay.addWidget(self.empty_lbl)

        self._rebuild()

    def set_templates(self, tpls):
        self._templates = tpls or []
        self._rebuild()

    def set_quick_names(self, names):
        valid = {t.get("name", "") for t in self._templates}
        self._quick_names = [n for n in (names or []) if n in valid]
        self._rebuild()

    def _rebuild(self):
        while self.items_layout.count() > 1:
            item = self.items_layout.takeAt(0)
            w = item.widget()
            if w:
                w.deleteLater()

        tpl_map = {t.get("name", ""): t for t in self._templates}
        shown = [n for n in self._quick_names if n in tpl_map]

        if not shown:
            self.scroll.setVisible(False)
            self.empty_lbl.setVisible(True)
            self.scroll.setFixedHeight(0)
            self._drag.setup([], None)
            return

        self.empty_lbl.setVisible(False)
        self.scroll.setVisible(True)

        qa_items = []
        for name in shown:
            tpl = tpl_map[name]
            it = QuickAccessItem(name, drag_controller=self._drag)
            it.clicked.connect(lambda _=None, t=tpl: self.template_selected.emit(t))
            it.removed.connect(lambda _=None, n=name: self._remove_name(n))
            self.items_layout.insertWidget(self.items_layout.count() - 1, it)
            qa_items.append(it)

        self._drag.setup(qa_items, self.items_layout)

        total_h = len(shown) * 44 + 4
        self.scroll.setFixedHeight(min(total_h, 176))

    def _remove_name(self, name):
        cur = list(SETTINGS.get("quick_access", []))
        if name in cur:
            cur.remove(name)
        SETTINGS["quick_access"] = cur
        save_settings()
        self._quick_names = [n for n in self._quick_names if n != name]
        self._rebuild()
        self.quick_access_changed.emit() 

    def _on_reorder(self, items):
        new_order = [getattr(it, "_name", "") for it in items
                     if getattr(it, "_name", "")]
        if not new_order:
            return
        SETTINGS["quick_access"] = new_order
        save_settings()
        self._quick_names = new_order
        self.quick_access_changed.emit() 


# ═════════════════════════════════════════════════════════════════
#  DropZone
# ═════════════════════════════════════════════════════════════════

class DropZone(QFrame):
    file_selected = Signal(str)

    def __init__(self):
        super().__init__()
        self.setAcceptDrops(True)
        self.setCursor(Qt.PointingHandCursor)
        self.setObjectName("dropzone")
        self._file_path = None
        self.setMinimumHeight(180)
        self._apply_style(active=False)

        lay = QVBoxLayout(self)
        lay.setContentsMargins(24, 30, 24, 30); lay.setSpacing(14)

        self.icon_lbl = QLabel()
        self.icon_lbl.setPixmap(icon_upload(56).pixmap(QSize(56, 56)))
        self.icon_lbl.setAlignment(Qt.AlignCenter)
        self.icon_lbl.setStyleSheet("background: transparent; border: none;")
        self.icon_lbl.setAttribute(Qt.WA_TransparentForMouseEvents)

        self.text_lbl = QLabel(tr(T.drop_zone_hint))
        self.text_lbl.setAlignment(Qt.AlignCenter)
        f = QFont(); f.setPointSize(11)
        self.text_lbl.setFont(f)
        self.text_lbl.setStyleSheet(f"color: {C_MUTED}; background: transparent; border: none;")
        self.text_lbl.setAttribute(Qt.WA_TransparentForMouseEvents)

        lay.addStretch(); lay.addWidget(self.icon_lbl)
        lay.addWidget(self.text_lbl); lay.addStretch()

    def _apply_style(self, active=False):
        border_color = C_GREEN if active else C_BORDER
        bg = C_LIGHT_GREEN if active else "#FAFCFA"
        self.setStyleSheet(f"""
            QFrame#dropzone {{
                background-color: {bg};
                border: 2px dashed {border_color};
                border-radius: 14px;
            }}
            QFrame#dropzone:hover {{
                border-color: {C_GREEN};
                background-color: {C_HOVER};
            }}
        """)

    def dragEnterEvent(self, e):
        if e.mimeData().hasUrls():
            urls = e.mimeData().urls()
            if urls and urls[0].toLocalFile().lower().endswith(".docx"):
                e.acceptProposedAction()
                self._apply_style(active=True); return
        e.ignore()

    def dragLeaveEvent(self, e):
        self._apply_style(active=False)

    def dropEvent(self, e):
        self._apply_style(active=False)
        urls = e.mimeData().urls()
        if urls:
            path = urls[0].toLocalFile()
            if path.lower().endswith(".docx"):
                self._set_file(path); e.acceptProposedAction()

    def mousePressEvent(self, e):
        if e.button() == Qt.LeftButton:
            start_dir = str(self._file_path and Path(self._file_path).parent or DOCUMENTS_DIR)
            path, _ = QFileDialog.getOpenFileName(
                self, tr(T.select_word_file), start_dir, tr(T.word_file_filter))
            if path:
                self._set_file(path)

    def _set_file(self, path):
        self._file_path = path
        self.icon_lbl.setPixmap(icon_check(56).pixmap(QSize(56, 56)))
        self.text_lbl.setText(Path(path).name)
        f = QFont(); f.setPointSize(11); f.setBold(True)
        self.text_lbl.setFont(f)
        self.text_lbl.setStyleSheet(f"color: {C_DARK_GREEN}; background: transparent; border: none;")
        self.file_selected.emit(path)

    def file_path(self):
        return self._file_path

    def reset(self):
        self._file_path = None
        self.icon_lbl.setPixmap(icon_upload(56).pixmap(QSize(56, 56)))
        self.text_lbl.setText(tr(T.drop_zone_hint))
        f = QFont(); f.setPointSize(11)
        self.text_lbl.setFont(f)
        self.text_lbl.setStyleSheet(f"color: {C_MUTED}; background: transparent; border: none;")
        self._apply_style(active=False)


class ExcelDropZone(QFrame):
    file_selected = Signal(str)
    SUPPORTED_SUFFIXES = (".xlsx", ".xlsm")

    def __init__(self):
        super().__init__()
        self.setAcceptDrops(True)
        self.setCursor(Qt.PointingHandCursor)
        self.setObjectName("excelDropZone")
        self._file_path = None
        self.setMinimumHeight(130)
        self._apply_style(active=False)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(18, 18, 18, 18)
        layout.setSpacing(8)
        self.icon_lbl = QLabel()
        self.icon_lbl.setPixmap(icon_upload(42).pixmap(QSize(42, 42)))
        self.icon_lbl.setAlignment(Qt.AlignCenter)
        self.icon_lbl.setStyleSheet("background: transparent; border: none;")
        self.icon_lbl.setAttribute(Qt.WA_TransparentForMouseEvents)
        self.text_lbl = QLabel(tr(T.batch_drop_hint))
        self.text_lbl.setAlignment(Qt.AlignCenter)
        self.text_lbl.setWordWrap(True)
        self.text_lbl.setStyleSheet(
            f"color: {C_MUTED}; background: transparent; border: none;"
        )
        self.text_lbl.setAttribute(Qt.WA_TransparentForMouseEvents)
        layout.addStretch()
        layout.addWidget(self.icon_lbl)
        layout.addWidget(self.text_lbl)
        layout.addStretch()

    def _apply_style(self, active=False):
        color = C_GREEN if active else C_BORDER
        background = C_LIGHT_GREEN if active else "#FAFCFA"
        self.setStyleSheet(f"""
            QFrame#excelDropZone {{
                background-color: {background};
                border: 2px dashed {color};
                border-radius: 12px;
            }}
            QFrame#excelDropZone:hover {{
                border-color: {C_GREEN}; background-color: {C_HOVER};
            }}
        """)

    @classmethod
    def _is_supported(cls, path):
        return Path(path).suffix.lower() in cls.SUPPORTED_SUFFIXES

    def dragEnterEvent(self, event):
        if event.mimeData().hasUrls():
            urls = event.mimeData().urls()
            if urls and self._is_supported(urls[0].toLocalFile()):
                event.acceptProposedAction()
                self._apply_style(active=True)
                return
        event.ignore()

    def dragLeaveEvent(self, event):
        self._apply_style(active=False)

    def dropEvent(self, event):
        self._apply_style(active=False)
        urls = event.mimeData().urls()
        if urls:
            path = urls[0].toLocalFile()
            if self._is_supported(path):
                self._set_file(path)
                event.acceptProposedAction()

    def mousePressEvent(self, event):
        if event.button() != Qt.LeftButton:
            return
        start_dir = str(
            Path(self._file_path).parent if self._file_path else DOCUMENTS_DIR
        )
        path, _ = QFileDialog.getOpenFileName(
            self,
            tr(T.batch_issue_title),
            start_dir,
            tr(T.batch_excel_filter),
        )
        if path:
            self._set_file(path)

    def _set_file(self, path):
        self._file_path = str(path)
        self.icon_lbl.setPixmap(icon_check(42).pixmap(QSize(42, 42)))
        self.text_lbl.setText(Path(path).name)
        font = QFont(); font.setPointSize(10); font.setBold(True)
        self.text_lbl.setFont(font)
        self.text_lbl.setStyleSheet(
            f"color: {C_DARK_GREEN}; background: transparent; border: none;"
        )
        self.file_selected.emit(self._file_path)

    def file_path(self):
        return self._file_path

    def reset(self):
        self._file_path = None
        self.icon_lbl.setPixmap(icon_upload(42).pixmap(QSize(42, 42)))
        self.text_lbl.setText(tr(T.batch_drop_hint))
        font = QFont(); font.setPointSize(10)
        self.text_lbl.setFont(font)
        self.text_lbl.setStyleSheet(
            f"color: {C_MUTED}; background: transparent; border: none;"
        )
        self._apply_style(active=False)


# ═════════════════════════════════════════════════════════════════
#  Small components
# ═════════════════════════════════════════════════════════════════

class RoundToggle(QPushButton):
    def __init__(self, label, tooltip="", bold=False, italic=False):
        super().__init__(label)
        self.setCheckable(True)
        self.setFixedSize(32, 32)
        self.setCursor(Qt.PointingHandCursor)
        self.setToolTip(tooltip)

        f = QFont()
        f.setPointSize(11)
        f.setBold(bold)
        f.setItalic(italic)
        self.setFont(f)

        self.setStyleSheet(f"""
            QPushButton {{
                background-color: #FFFFFF;
                border: 1.5px solid {C_BORDER};
                border-radius: 16px;
                color: {C_MUTED};
            }}
            QPushButton:hover {{
                border-color: {C_GREEN};
                color: {C_DARK_GREEN};
            }}
            QPushButton:checked {{
                background-color: {C_GREEN};
                border-color: {C_GREEN};
                color: white;
            }}
        """)


FONT_SIZES = ["Auto", "8", "9", "10", "11", "12", "14", "16",
              "18", "20", "22", "24", "26", "28", "36", "48", "72"]


class FontSizeCombo(ComboWithArrow):
    def __init__(self):
        super().__init__(editable=True)
        self.setFixedHeight(32)
        self.setMinimumWidth(100)
        self.addItems([tr(T.auto)] + FONT_SIZES[1:])
        self.setCurrentText(tr(T.auto))
        self.setToolTip(tr(T.font_size_tooltip))

        self.setStyleSheet(f"""
            QComboBox {{
                background-color: #FFFFFF;
                border: 1.5px solid {C_BORDER};
                border-radius: 8px;
                color: {C_DARK_GREEN};
                font-size: 11px;
            }}
            QComboBox:focus {{ border-color: {C_GREEN}; }}
            QComboBox QAbstractItemView {{
                background-color: #FFFFFF;
                border: 1px solid {C_BORDER};
                border-radius: 6px;
                padding: 2px;
                outline: none;
                color: {C_DARK_GREEN};
                selection-background-color: {C_LIGHT_GREEN};
                selection-color: {C_DARK_GREEN};
                max-height: 220px;
            }}
            QComboBox QAbstractItemView::item {{
                padding: 4px 8px; border-radius: 4px; min-height: 20px;
            }}
            QComboBox QAbstractItemView::item:hover {{
                background-color: {C_HOVER};
            }}
        """)

        le = self.lineEdit()
        if le is not None:
            size_pattern = QRegularExpression(
                r"(?:[0-9۰-۹٠-٩]+(?:[.,٫][0-9۰-۹٠-٩]*)?|[.,٫][0-9۰-۹٠-٩]+)?"
            )
            le.setValidator(QRegularExpressionValidator(size_pattern, le))
            le.editingFinished.connect(self._ensure_not_empty)

    def _ensure_not_empty(self):
        le = self.lineEdit()
        if le is None:
            return
        if not le.text().strip():
            le.setText(tr(T.auto))

    def eventFilter(self, obj, event):
        if super().eventFilter(obj, event):
            return True

        le = self.lineEdit()
        if obj is le:
            et = event.type()
            if et == QEvent.MouseButtonPress:
                if le.text().strip():
                    QTimer.singleShot(0, le.clear)
            elif et == QEvent.FocusOut:
                QTimer.singleShot(0, self._ensure_not_empty)
        return False


# ═════════════════════════════════════════════════════════════════
#  HomePage
# ═════════════════════════════════════════════════════════════════

class HomePage(QWidget):
    settings_requested = Signal()
    quick_template_selected = Signal(dict)
    quick_access_changed = Signal()  

    def __init__(self):
        super().__init__()
        self.setStyleSheet("background: transparent;")
        self._templates = []

        outer = QVBoxLayout(self)
        outer.setContentsMargins(40, 18, 40, 24)
        outer.setSpacing(0)

        top_row = QHBoxLayout()
        gear_btn = QPushButton(tr(T.settings_spaced))
        gear_btn.setIcon(icon_gear(20))
        gear_btn.setIconSize(QSize(20, 20))
        gear_btn.setLayoutDirection(_dir())
        gear_btn.setFixedHeight(40)
        gear_btn.setMinimumWidth(120)
        gear_btn.setCursor(Qt.PointingHandCursor)
        gear_btn.setToolTip(tr(T.settings))
        gear_btn.setStyleSheet(f"""
            QPushButton {{
                background-color: transparent;
                color: {C_DARK_GREEN};
                border: 1px solid transparent;
                border-radius: 10px;
                font-size: 11px;
                font-weight: bold;
                padding: 0 14px;
            }}
            QPushButton:hover {{
                background-color: {C_LIGHT_GREEN};
                border-color: {C_BORDER};
            }}
        """)
        gear_btn.clicked.connect(self.settings_requested.emit)

        guide_btn = QPushButton(tr(T.guide))
        guide_btn.setIcon(icon_help(19))
        guide_btn.setIconSize(QSize(19, 19))
        guide_btn.setLayoutDirection(_dir())
        guide_btn.setFixedHeight(38)
        guide_btn.setMinimumWidth(120)
        guide_btn.setCursor(Qt.PointingHandCursor)
        guide_btn.setToolTip(tr(T.guide_tooltip))
        guide_btn.setStyleSheet(f"""
            QPushButton {{
                background-color: transparent;
                color: {C_DARK_GREEN};
                border: 1px solid transparent;
                border-radius: 10px;
                font-size: 11px;
                font-weight: bold;
                padding: 0 14px;
            }}
            QPushButton:hover {{
                background-color: {C_LIGHT_GREEN};
                border-color: {C_BORDER};
            }}
        """)
        guide_btn.clicked.connect(self._open_guide)

        navigation = QVBoxLayout()
        navigation.setContentsMargins(0, 0, 0, 0)
        navigation.setSpacing(2)
        navigation.addWidget(gear_btn)
        navigation.addWidget(guide_btn)
        top_row.addLayout(navigation)
        top_row.addStretch()
        outer.addLayout(top_row)

        outer.addSpacing(4)

        # Scroll container for the whole main area (small screens)
        main_scroll = QScrollArea()
        main_scroll.setWidgetResizable(True)
        main_scroll.setFrameShape(QFrame.NoFrame)
        main_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        main_scroll.setStyleSheet(f"""
            QScrollArea {{ background: transparent; border: none; }}
            QScrollBar:vertical {{
                background: transparent; width: 6px;
            }}
            QScrollBar::handle:vertical {{
                background: {C_BORDER}; border-radius: 3px; min-height: 30px;
            }}
            QScrollBar::handle:vertical:hover {{ background: {C_GREEN}; }}
            QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical {{
                height: 0;
            }}
        """)

        main_content = QWidget()
        main_content.setStyleSheet("background: transparent;")
        mc_lay = QVBoxLayout(main_content)
        mc_lay.setContentsMargins(0, 0, 0, 0)
        mc_lay.setSpacing(18)

        card = QWidget(); card.setStyleSheet("background: transparent;")
        card.setFixedWidth(440)
        cl = QVBoxLayout(card)
        cl.setContentsMargins(0, 0, 0, 0); cl.setSpacing(8)

        big = QLabel()
        big.setPixmap(icon_document_big(72).pixmap(QSize(72, 72)))
        big.setAlignment(Qt.AlignCenter)
        big.setStyleSheet("background: transparent;")

        title = QLabel(tr(T.issue_letter))
        tf = QFont(); tf.setPointSize(22); tf.setBold(True)
        title.setFont(tf); title.setAlignment(Qt.AlignCenter)
        title.setStyleSheet(f"color: {C_DARK_GREEN}; background: transparent;")

        self.dropdown = TemplateDropdown()

        cl.addWidget(big, 0, Qt.AlignHCenter)
        cl.addWidget(title); cl.addSpacing(12); cl.addWidget(self.dropdown)

        center = QHBoxLayout()
        center.addStretch(); center.addWidget(card); center.addStretch()
        mc_lay.addLayout(center)

        qa_wrap = QWidget()
        qa_wrap.setStyleSheet("background: transparent;")
        qa_wrap.setFixedWidth(440)
        qa_lay = QVBoxLayout(qa_wrap)
        qa_lay.setContentsMargins(0, 0, 0, 0)

        self.quick_access = QuickAccessWidget()
        self.quick_access.template_selected.connect(self.quick_template_selected.emit)
        self.quick_access.add_requested.connect(self._on_qa_add)
        qa_lay.addWidget(self.quick_access)
        self.quick_access.quick_access_changed.connect(self.quick_access_changed.emit)
        center_qa = QHBoxLayout()
        center_qa.addStretch(); center_qa.addWidget(qa_wrap); center_qa.addStretch()
        mc_lay.addLayout(center_qa)

        mc_lay.addStretch(1)

        main_scroll.setWidget(main_content)
        outer.addWidget(main_scroll, 1)

    def set_templates(self, tpls):
        self._templates = tpls or []
        self.dropdown.set_templates(self._templates)
        self.quick_access.set_templates(self._templates)
        self.quick_access.set_quick_names(SETTINGS.get("quick_access", []))

    def _on_qa_add(self):
        valid_names = {t.get("name", "") for t in self._templates}
        current = [n for n in SETTINGS.get("quick_access", [])
                   if n in valid_names]
        dlg = QuickAccessPicker(self, self._templates, current)
        if dlg.exec() == QDialog.Accepted:
            new_names = dlg.selected_names()
            SETTINGS["quick_access"] = new_names
            save_settings()
            self.quick_access.set_quick_names(new_names)
            self.quick_access_changed.emit()      # ← جدید

    def _open_guide(self):
        language = SETTINGS.get("language", "fa")
        video_name = (
            "English_guide.mp4" if language == "en" else "Persian_guide.mp4"
        )
        program_dir = (
            Path(sys.executable).resolve().parent
            if getattr(sys, "frozen", False)
            else Path(__file__).resolve().parent
        )
        video_path = program_dir / "Guide" / video_name

        if not video_path.is_file():
            ThemedDialog.warning(
                self, tr(T.error), tr(T.guide_not_found, path=video_path)
            )
            return

        try:
            video_url = QUrl.fromLocalFile(str(video_path.resolve()))
            if not QDesktopServices.openUrl(video_url):
                raise OSError("The default video player did not accept the file.")
        except Exception as exc:
            ThemedDialog.error(
                self, tr(T.error), tr(T.guide_open_error, err=exc)
            )


# ═════════════════════════════════════════════════════════════════
#  AddLetterPage
# ═════════════════════════════════════════════════════════════════

class AddLetterPage(QWidget):
    back_requested = Signal()
    saved = Signal()

    def __init__(self):
        super().__init__()
        self.setStyleSheet("background: transparent;")

        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0); root.setSpacing(0)

        # ─ Header (static) ─
        header_wrap = QWidget()
        header_wrap.setStyleSheet("background: transparent;")
        hw_lay = QHBoxLayout(header_wrap)
        hw_lay.setContentsMargins(40, 22, 40, 14); hw_lay.setSpacing(10)

        title = QLabel(tr(T.add_new_letter))
        tf = QFont(); tf.setPointSize(15); tf.setBold(True)
        title.setFont(tf)
        title.setStyleSheet(f"color: {C_DARK_GREEN}; background: transparent;")

        back_btn = QPushButton(tr(T.back_spaced))
        back_btn.setIcon(icon_arrow_left(16)); back_btn.setIconSize(QSize(16, 16))
        back_btn.setCursor(Qt.PointingHandCursor); back_btn.setFixedHeight(34)
        back_btn.setStyleSheet(f"""
            QPushButton {{
                background: transparent; border: 1px solid {C_BORDER};
                border-radius: 8px; color: {C_DARK_GREEN};
                padding: 0 14px; font-size: 11px;
            }}
            QPushButton:hover {{
                background-color: {C_LIGHT_GREEN};
                border-color: {C_GREEN};
            }}
        """)
        back_btn.clicked.connect(self.back_requested.emit)

        hw_lay.addWidget(title); hw_lay.addStretch(); hw_lay.addWidget(back_btn)
        root.addWidget(header_wrap)

        # ─ Scrollable body ─
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.NoFrame)
        scroll.setStyleSheet(f"""
            QScrollArea {{ background: transparent; border: none; }}
            QScrollBar:vertical {{
                background: transparent; width: 8px;
            }}
            QScrollBar::handle:vertical {{
                background: {C_BORDER}; border-radius: 4px; min-height: 30px;
            }}
            QScrollBar::handle:vertical:hover {{ background: {C_GREEN}; }}
            QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical {{
                height: 0;
            }}
        """)

        body = QWidget()
        body.setStyleSheet("background: transparent;")
        bl = QVBoxLayout(body)
        bl.setContentsMargins(40, 0, 40, 20)
        bl.setSpacing(14)

        nf = QFont(); nf.setPointSize(10); nf.setBold(True)

        name_lbl = QLabel(tr(T.letter_name))
        name_lbl.setFont(nf)
        name_lbl.setStyleSheet(f"color: {C_DARK_GREEN}; background: transparent;")
        bl.addWidget(name_lbl)

        self.name_input = QLineEdit()
        self.name_input.setPlaceholderText(tr(T.ph_letter_name))
        self.name_input.setFixedHeight(44)
        self.name_input.setStyleSheet(f"""
            QLineEdit {{
                background-color: #FFFFFF;
                border: 1.5px solid {C_BORDER};
                border-radius: 10px;
                padding: 0 14px;
                color: {C_DARK_GREEN};
                font-size: 11px;
            }}
            QLineEdit:focus {{ border-color: {C_GREEN}; }}
        """)
        bl.addWidget(self.name_input)

        drop_lbl = QLabel(tr(T.word_template_file))
        drop_lbl.setFont(nf)
        drop_lbl.setStyleSheet(f"color: {C_DARK_GREEN}; background: transparent;")
        bl.addWidget(drop_lbl)

        self.drop_zone = DropZone()
        bl.addWidget(self.drop_zone)

        bl.addStretch(1)

        scroll.setWidget(body)
        root.addWidget(scroll, 1)

        # ─ Footer (static) ─
        footer_wrap = QWidget()
        footer_wrap.setStyleSheet("background: transparent;")
        fw_lay = QVBoxLayout(footer_wrap)
        fw_lay.setContentsMargins(40, 8, 40, 22)

        self.confirm_btn = QPushButton(tr(T.confirm_save))
        self.confirm_btn.setCursor(Qt.PointingHandCursor)
        self.confirm_btn.setFixedHeight(48)
        self.confirm_btn.setStyleSheet(f"""
            QPushButton {{
                background-color: {C_GREEN}; color: white;
                border: none; border-radius: 10px;
                font-size: 12px; font-weight: bold;
            }}
            QPushButton:hover {{ background-color: {C_DARK_GREEN}; }}
        """)
        self.confirm_btn.clicked.connect(self._on_confirm)
        fw_lay.addWidget(self.confirm_btn)
        root.addWidget(footer_wrap)

    def reset(self):
        self.name_input.clear(); self.drop_zone.reset()

    def _on_confirm(self):
        if not HAS_DOCX:
            ThemedDialog.error(self, tr(T.missing_library), tr(T.no_docx))
            return

        name = self.name_input.text().strip()
        file_path = self.drop_zone.file_path()

        if not name:
            ThemedDialog.warning(self, tr(T.error), tr(T.err_empty_name))
            self.name_input.setFocus(); return
        if not file_path:
            ThemedDialog.warning(self, tr(T.error), tr(T.err_empty_file))
            return

        existing = load_templates_raw()
        if any(t.get("name") == name for t in existing):
            ThemedDialog.warning(self, tr(T.error),
                                 tr(T.err_name_exists, name=name))
            return

        dst = None
        try:
            WORDS_DIR.mkdir(parents=True, exist_ok=True)
            src = Path(file_path)
            dst = WORDS_DIR / src.name
            counter = 1
            while dst.exists():
                dst = WORDS_DIR / f"{src.stem}_{counter}{src.suffix}"
                counter += 1
            copy2(str(src), str(dst))

            try:
                props = extract_placeholders(str(dst))
            except WrapperCountError as we:
                try:
                    dst.unlink()
                except Exception:
                    pass
                ThemedDialog.error(self, tr(T.err_wrapper_count), str(we))
                return

            # ─── هیچ متغیری در فایل پیدا نشد → ذخیره نکن ───
            if not props:
                try:
                    dst.unlink()
                except Exception:
                    pass
                ThemedDialog.warning(
                    self, tr(T.error),
                    tr(T.err_no_vars, w=_get_wrapper()))
                return

            new_obj = {"name": name, "address": str(dst), "Properties": props}
            existing.append(new_obj)
            save_templates(existing)

            ThemedDialog.info(self, tr(T.saved),tr(T.saved_with_n, name=name, n=len(props)))

            self.reset(); self.saved.emit()
        except Exception as ex:
            if dst is not None and dst.exists():
                try:
                    dst.unlink()
                except Exception:
                    pass
            ThemedDialog.error(self, tr(T.error),tr(T.err_generic, err=ex))
            print("[error]", ex)


# ═════════════════════════════════════════════════════════════════
#  EditLetterPage — header static, everything else scrolls
# ═════════════════════════════════════════════════════════════════

class EditLetterPage(QWidget):
    back_requested = Signal()
    batch_requested = Signal()

    def __init__(self):
        super().__init__()
        self.setStyleSheet("background: transparent;")
        self._template = None
        self._fields = {}

        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0); root.setSpacing(0)

        # ─ Header (static) ─
        header_wrap = QWidget()
        header_wrap.setStyleSheet("background: transparent;")
        hw_lay = QHBoxLayout(header_wrap)
        hw_lay.setContentsMargins(40, 22, 40, 14); hw_lay.setSpacing(10)

        self.title_lbl = QLabel(tr(T.issue_letter))
        tf = QFont(); tf.setPointSize(15); tf.setBold(True)
        self.title_lbl.setFont(tf)
        self.title_lbl.setStyleSheet(f"color: {C_DARK_GREEN}; background: transparent;")

        back_btn = QPushButton(tr(T.back_spaced))
        back_btn.setIcon(icon_arrow_left(16)); back_btn.setIconSize(QSize(16, 16))
        back_btn.setCursor(Qt.PointingHandCursor); back_btn.setFixedHeight(34)
        back_btn.setStyleSheet(f"""
            QPushButton {{
                background: transparent; border: 1px solid {C_BORDER};
                border-radius: 8px; color: {C_DARK_GREEN};
                padding: 0 14px; font-size: 11px;
            }}
            QPushButton:hover {{
                background-color: {C_LIGHT_GREEN};
                border-color: {C_GREEN};
            }}
        """)
        back_btn.clicked.connect(self.back_requested.emit)

        self.batch_btn = QPushButton(tr(T.batch_excel_issue))
        self.batch_btn.setIcon(icon_document(17))
        self.batch_btn.setIconSize(QSize(17, 17))
        self.batch_btn.setCursor(Qt.PointingHandCursor)
        self.batch_btn.setFixedHeight(34)
        self.batch_btn.setToolTip(tr(T.batch_excel_issue))
        self.batch_btn.setStyleSheet(f"""
            QPushButton {{
                background-color: {C_LIGHT_GREEN};
                border: 1px solid {C_BORDER};
                border-radius: 8px;
                color: {C_DARK_GREEN};
                font-size: 10px; font-weight: bold;
                padding: 0 12px;
            }}
            QPushButton:hover {{
                background-color: {C_HOVER}; border-color: {C_GREEN};
            }}
        """)
        self.batch_btn.clicked.connect(self.batch_requested.emit)

        hw_lay.addWidget(self.title_lbl)
        hw_lay.addStretch()
        hw_lay.addWidget(self.batch_btn)
        hw_lay.addWidget(back_btn)
        root.addWidget(header_wrap)

        # ─ Everything else in a scroll ─
        self.scroll = QScrollArea()
        self.scroll.setWidgetResizable(True)
        self.scroll.setFrameShape(QFrame.NoFrame)
        self.scroll.setStyleSheet(f"""
            QScrollArea {{ background: transparent; border: none; }}
            QScrollBar:vertical {{
                background: transparent; width: 8px;
            }}
            QScrollBar::handle:vertical {{
                background: {C_BORDER}; border-radius: 4px; min-height: 30px;
            }}
            QScrollBar::handle:vertical:hover {{ background: {C_GREEN}; }}
            QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical {{
                height: 0;
            }}
        """)

        self.scroll_body = QWidget()
        self.scroll_body.setStyleSheet("background: transparent;")
        sb = QVBoxLayout(self.scroll_body)
        sb.setContentsMargins(40, 0, 40, 20)
        sb.setSpacing(14)

        nf = QFont(); nf.setPointSize(10); nf.setBold(True)

        fn_lbl = QLabel(tr(T.output_filename))
        fn_lbl.setFont(nf)
        fn_lbl.setStyleSheet(f"color: {C_DARK_GREEN}; background: transparent;")
        sb.addWidget(fn_lbl)

        self.filename_input = QLineEdit()
        self.filename_input.setPlaceholderText(tr(T.ph_filename))
        self.filename_input.setFixedHeight(44)
        self.filename_input.setStyleSheet(f"""
            QLineEdit {{
                background-color: #FFFFFF;
                border: 1.5px solid {C_BORDER};
                border-radius: 10px;
                padding: 0 14px;
                color: {C_DARK_GREEN};
                font-size: 11px;
            }}
            QLineEdit:focus {{ border-color: {C_GREEN}; }}
        """)
        sb.addWidget(self.filename_input)

        vars_lbl = QLabel(tr(T.variable_values))
        vars_lbl.setFont(nf)
        vars_lbl.setStyleSheet(f"color: {C_DARK_GREEN}; background: transparent;")
        sb.addWidget(vars_lbl)

        self.fields_container = QWidget()
        self.fields_container.setStyleSheet("background: transparent;")
        self.fields_layout = QVBoxLayout(self.fields_container)
        self.fields_layout.setContentsMargins(0, 0, 0, 0)
        self.fields_layout.setSpacing(12)
        self.fields_layout.addStretch()

        sb.addWidget(self.fields_container)
        sb.addStretch(1)

        self.scroll.setWidget(self.scroll_body)
        root.addWidget(self.scroll, 1)

        # ─ Footer (static): confirm button ─
        footer_wrap = QWidget()
        footer_wrap.setStyleSheet("background: transparent;")
        fw_lay = QHBoxLayout(footer_wrap)
        fw_lay.setContentsMargins(40, 8, 40, 22)

        self.confirm_btn = QPushButton(tr(T.confirm_issue))
        self.confirm_btn.setCursor(Qt.PointingHandCursor)
        self.confirm_btn.setFixedHeight(46)
        self.confirm_btn.setMinimumWidth(200)
        self.confirm_btn.setStyleSheet(f"""
            QPushButton {{
                background-color: {C_GREEN}; color: white;
                border: none; border-radius: 10px;
                font-size: 12px; font-weight: bold;
                padding: 0 28px;
            }}
            QPushButton:hover {{ background-color: {C_DARK_GREEN}; }}
        """)
        self.confirm_btn.clicked.connect(self._on_confirm)
        fw_lay.addWidget(self.confirm_btn)
        fw_lay.addStretch()
        root.addWidget(footer_wrap)

    def load_template(self, tpl):
        self._template = tpl
        self.title_lbl.setText(tr(T.issue_named, name=tpl.get('name', '')))
        self.filename_input.setText(tpl.get("name", ""))

        self._fields.clear()
        while self.fields_layout.count() > 1:
            item = self.fields_layout.takeAt(0)
            w = item.widget()
            if w:
                w.deleteLater()

        props = tpl.get("Properties", {})

        normalized = {}
        for k, v in props.items():
            k2 = (k or "").strip()
            if not k2:
                continue
            if isinstance(v, dict):
                try:
                    start = int(v.get("start", 0))
                except Exception:
                    start = 0
                normalized[k2] = {
                    "start": start,
                    "auto_size": v.get("auto_size"),
                    "font_name": v.get("font_name"),
                }
            else:
                try:
                    start = int(v) if v is not None else 0
                except Exception:
                    start = 0
                normalized[k2] = {
                    "start": start,
                    "auto_size": None,
                    "font_name": None,
                }

        if not normalized:
            empty = QLabel(tr(T.no_variables))
            empty.setStyleSheet(f"color: {C_MUTED}; background: transparent; padding: 8px;")
            empty.setFont(QFont("", 10))
            self.fields_layout.insertWidget(0, empty)
            return

        sorted_items = sorted(normalized.items(),
                              key=lambda kv: kv[1]["start"])
        for idx, (key, info) in enumerate(sorted_items):
            box = self._build_field(
                key,
                show_labels=(idx == 0),
                default_size=None,
            )
            self.fields_layout.insertWidget(self.fields_layout.count() - 1, box)

    def _pick_color(self, key):
        widgets = self._fields.get(key)
        if not widgets:
            return
        btn = widgets.get("color")
        if btn is None:
            return

        cur = btn.property("color") or "#000000"
        dlg = QColorDialog(QColor(cur), self)
        dlg.setOption(QColorDialog.ColorDialogOption.DontUseNativeDialog, True)
        dlg.setOption(QColorDialog.ColorDialogOption.ShowAlphaChannel, False)
        # QColorDialog keeps a conventional left-to-right layout in both app languages.
        dlg.setLayoutDirection(Qt.LeftToRight)
        dlg.setWindowTitle(tr(T.pick_color))
        dlg.setAutoFillBackground(True)

        palette = dlg.palette()
        palette.setColor(QPalette.ColorRole.Window, QColor("#FFFFFF"))
        palette.setColor(QPalette.ColorRole.WindowText, QColor(C_DARK_GREEN))
        palette.setColor(QPalette.ColorRole.Base, QColor("#FFFFFF"))
        palette.setColor(QPalette.ColorRole.AlternateBase, QColor(C_BG))
        palette.setColor(QPalette.ColorRole.Text, QColor(C_DARK_GREEN))
        palette.setColor(QPalette.ColorRole.Button, QColor(C_LIGHT_GREEN))
        palette.setColor(QPalette.ColorRole.ButtonText, QColor(C_DARK_GREEN))
        palette.setColor(QPalette.ColorRole.Highlight, QColor(C_GREEN))
        palette.setColor(QPalette.ColorRole.HighlightedText, QColor("#FFFFFF"))
        dlg.setPalette(palette)

        dlg.setStyleSheet(f"""
            QColorDialog {{
                background-color: #FFFFFF;
                color: {C_DARK_GREEN};
            }}
            QColorDialog QLabel {{
                color: {C_DARK_GREEN};
                background-color: transparent;
            }}
            QColorDialog QLineEdit,
            QColorDialog QSpinBox,
            QColorDialog QDoubleSpinBox,
            QColorDialog QComboBox {{
                color: {C_DARK_GREEN};
                background-color: #FFFFFF;
                border: 1px solid {C_BORDER};
                border-radius: 4px;
                padding: 3px 6px;
            }}
            QColorDialog QPushButton {{
                color: {C_DARK_GREEN};
                background-color: {C_LIGHT_GREEN};
                border: 1px solid {C_BORDER};
                border-radius: 6px;
                padding: 5px 12px;
            }}
            QColorDialog QPushButton:hover {{
                background-color: {C_HOVER};
                border-color: {C_GREEN};
            }}
        """)

        def localize_color_dialog_labels():
            # Force all already-created child controls to LTR as well; the
            # parent application may use RTL for Persian pages.
            dlg.setLayoutDirection(Qt.LeftToRight)
            for control in dlg.findChildren(QWidget):
                control.setLayoutDirection(Qt.LeftToRight)
            for label in dlg.findChildren(QLabel):
                label_text = label.text().replace("&", "").strip().rstrip(":").strip().casefold()
                if label_text == "html":
                    label.setText("HEX:")
                    break
            for button in dlg.findChildren(QPushButton):
                button_text = button.text().replace("&", "").strip().casefold()
                if button_text == "ok":
                    button.setText(tr(T.ok))
                elif button_text == "cancel":
                    button.setText(tr(T.cancel))

        # Run after Qt has created the non-native dialog controls.
        QTimer.singleShot(0, localize_color_dialog_labels)

        if dlg.exec():
            col = dlg.currentColor()
            if col.isValid():
                hexc = col.name()
                btn.setProperty("color", hexc)
                apply_color_button_style(btn, hexc)

    def _build_field(self, key, show_labels=False, default_size=None):
        box = QWidget()
        box.setStyleSheet("background: transparent;")
        bl = QVBoxLayout(box)
        bl.setContentsMargins(0, 0, 0, 0); bl.setSpacing(6)

        top = QHBoxLayout()
        top.setSpacing(8)

        lbl = QLabel(key.strip())
        lf = QFont(); lf.setPointSize(10); lf.setBold(True)
        lbl.setFont(lf)
        lbl.setStyleSheet(f"color: {C_DARK_GREEN}; background: transparent;")
        top.addWidget(lbl)
        top.addStretch()

        if show_labels:
            size_lbl = QLabel(tr(T.font_size))
            size_lbl.setStyleSheet(
                f"color: {C_MUTED}; background: transparent; font-size: 10px;")
            top.addWidget(size_lbl)

        size_combo = FontSizeCombo()
        if default_size is not None:
            try:
                if isinstance(default_size, (int, float)):
                    size_combo.setCurrentText(f"{default_size:g}")
                else:
                    size_combo.setCurrentText(str(default_size))
            except Exception:
                pass
        top.addWidget(size_combo)

        # Bold
        bold_wrap = QWidget()
        bold_wrap.setStyleSheet("background: transparent;")
        bold_lay = QHBoxLayout(bold_wrap)
        bold_lay.setContentsMargins(0, 0, 0, 0); bold_lay.setSpacing(6)
        bold_btn = RoundToggle("B", tr(T.bold), bold=True)
        if show_labels:
            bold_lbl = QLabel(tr(T.bold))
            bold_lbl.setStyleSheet(
                f"color: {C_MUTED}; background: transparent; font-size: 10px;")
            bold_lbl.setCursor(Qt.PointingHandCursor)
            bold_lbl.mousePressEvent = lambda _e, b=bold_btn: b.toggle()
            bold_lay.addWidget(bold_lbl)
        bold_lay.addWidget(bold_btn)
        top.addWidget(bold_wrap)

        # Italic
        ital_wrap = QWidget()
        ital_wrap.setStyleSheet("background: transparent;")
        ital_lay = QHBoxLayout(ital_wrap)
        ital_lay.setContentsMargins(0, 0, 0, 0); ital_lay.setSpacing(6)
        italic_btn = RoundToggle("I", tr(T.italic), italic=True)
        if show_labels:
            italic_lbl = QLabel(tr(T.italic))
            italic_lbl.setStyleSheet(
                f"color: {C_MUTED}; background: transparent; font-size: 10px;")
            italic_lbl.setCursor(Qt.PointingHandCursor)
            italic_lbl.mousePressEvent = lambda _e, b=italic_btn: b.toggle()
            ital_lay.addWidget(italic_lbl)
        ital_lay.addWidget(italic_btn)
        top.addWidget(ital_wrap)

        # Color
        color_wrap = QWidget()
        color_wrap.setStyleSheet("background: transparent;")
        color_lay = QHBoxLayout(color_wrap)
        color_lay.setContentsMargins(0, 0, 0, 0); color_lay.setSpacing(6)
        if show_labels:
            color_lbl = QLabel(tr(T.color))
            color_lbl.setStyleSheet(
                f"color: {C_MUTED}; background: transparent; font-size: 10px;")
            color_lay.addWidget(color_lbl)
        color_btn = QPushButton()
        color_btn.setFixedSize(16, 16)
        color_btn.setCursor(Qt.PointingHandCursor)
        color_btn.setToolTip(tr(T.color_tooltip))
        color_btn.setProperty("color", "#000000")
        apply_color_button_style(color_btn, "#000000")
        color_btn.clicked.connect(lambda _=False, k=key: self._pick_color(k))
        color_lay.addWidget(color_btn)
        top.addWidget(color_wrap)

        bl.addLayout(top)

        inp = QLineEdit()
        inp.setFixedHeight(42)
        inp.setStyleSheet(f"""
            QLineEdit {{
                background-color: #FFFFFF;
                border: 1.5px solid {C_BORDER};
                border-radius: 10px;
                padding: 0 14px;
                color: {C_DARK_GREEN};
                font-size: 11px;
            }}
            QLineEdit:focus {{ border-color: {C_GREEN}; }}
        """)
        bl.addWidget(inp)

        self._fields[key] = {
            "input": inp,
            "size": size_combo,
            "bold": bold_btn,
            "italic": italic_btn,
            "color": color_btn,
        }
        return box

    def _on_confirm(self):
        if not HAS_DOCX:
            ThemedDialog.error(self, tr(T.missing_library), tr(T.no_docx))
            return
        if not self._template:
            return

        filename = self.filename_input.text().strip()
        if not filename:
            ThemedDialog.warning(self, tr(T.error), tr(T.err_empty_filename))
            return

        if self._fields:
            any_filled = any(
                w["input"].text().strip() for w in self._fields.values()
            )
            if not any_filled:
                ThemedDialog.warning(self, tr(T.error), tr(T.err_fill_one))
                return

        for ch in '<>:"/\\|?*':
            filename = filename.replace(ch, "_")
        filename = filename.strip()
        if not filename.lower().endswith(".docx"):
            filename += ".docx"

        output_dir = Path(SETTINGS.get("output_dir", str(Path.home() / "Desktop")))
        try:
            output_dir.mkdir(parents=True, exist_ok=True)
        except Exception as e:
            ThemedDialog.error(self, tr(T.error),
                               tr(T.err_output_dir_create, err=e))
            return

        out_path = output_dir / filename
        src = Path(self._template.get("address", ""))
        if not src.exists():
            ThemedDialog.error(self, tr(T.error),
                               tr(T.err_template_not_found, path=src))
            return

        if out_path.exists():
            ok = ThemedDialog.confirm(
                self, tr(T.file_exists),
                tr(T.file_overwrite, f=filename))
            if not ok: return

        replacements = {}
        font_overrides = {}
        for key, widgets in self._fields.items():
            replacements[key] = widgets["input"].text()

            size_text = widgets["size"].currentText().strip().lower()
            if size_text in ("", "auto", tr(T.auto).strip().lower()):
                size_val = None
            else:
                try:
                    digit_map = str.maketrans(
                        "۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩",
                        "01234567890123456789",
                    )
                    normalized_size = (
                        size_text.translate(digit_map)
                        .replace(",", ".")
                        .replace("٫", ".")
                    )
                    size_val = float(normalized_size)
                except ValueError:
                    size_val = None
            bold_val = widgets["bold"].isChecked()
            italic_val = widgets["italic"].isChecked()
            color_val = widgets["color"].property("color") or "#000000"
            font_overrides[key] = (size_val, bold_val, italic_val, color_val)

        try:
            replace_placeholders_in_docx(
                str(src), str(out_path), replacements, font_overrides)
            ThemedDialog.info(
                self, tr(T.issued),
                tr(T.issued_ok, path=out_path))
            self.back_requested.emit()
        except Exception as ex:
            ThemedDialog.error(self, tr(T.error), tr(T.err_issue, err=ex))
            print("[error]", ex)


class BulkIssuePage(QWidget):
    back_requested = Signal()

    def __init__(self):
        super().__init__()
        self._template = None
        self.variable_names = []
        self._font_options = {}
        self.setStyleSheet("background: transparent;")

        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(0)

        header = QWidget()
        header.setStyleSheet("background: transparent;")
        header_layout = QHBoxLayout(header)
        header_layout.setContentsMargins(40, 18, 40, 12)
        header_layout.setSpacing(10)

        title = QLabel(tr(T.batch_issue_title))
        title_font = QFont(); title_font.setPointSize(15); title_font.setBold(True)
        title.setFont(title_font)
        title.setStyleSheet(f"color: {C_DARK_GREEN}; background: transparent;")

        back_btn = QPushButton(tr(T.back_spaced))
        back_btn.setIcon(icon_arrow_left(16))
        back_btn.setIconSize(QSize(16, 16))
        back_btn.setCursor(Qt.PointingHandCursor)
        back_btn.setFixedHeight(34)
        back_btn.setStyleSheet(f"""
            QPushButton {{
                background: transparent; border: 1px solid {C_BORDER};
                border-radius: 8px; color: {C_DARK_GREEN};
                padding: 0 14px; font-size: 11px;
            }}
            QPushButton:hover {{
                background-color: {C_LIGHT_GREEN}; border-color: {C_GREEN};
            }}
        """)
        back_btn.clicked.connect(self.back_requested.emit)
        header_layout.addWidget(title)
        header_layout.addStretch()
        header_layout.addWidget(back_btn)
        root.addWidget(header)

        self.scroll = QScrollArea()
        self.scroll.setWidgetResizable(True)
        self.scroll.setFrameShape(QFrame.NoFrame)
        self.scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.scroll.setStyleSheet(f"""
            QScrollArea {{ background: transparent; border: none; }}
            QScrollBar:vertical {{
                background: transparent; width: 8px;
            }}
            QScrollBar::handle:vertical {{
                background: {C_BORDER}; border-radius: 4px; min-height: 30px;
            }}
            QScrollBar::handle:vertical:hover {{ background: {C_GREEN}; }}
            QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical {{
                height: 0;
            }}
        """)

        body = QWidget()
        body.setStyleSheet("background: transparent;")
        body_layout = QVBoxLayout(body)
        body_layout.setContentsMargins(40, 0, 40, 18)
        body_layout.setSpacing(10)

        self.explanation_lbl = QLabel(tr(T.batch_explanation))
        self.explanation_lbl.setWordWrap(True)
        self.explanation_lbl.setStyleSheet(
            f"color: {C_DARK_GREEN}; background: {C_LIGHT_GREEN}; "
            f"border: 1px solid {C_BORDER}; border-radius: 10px; padding: 10px;"
        )
        body_layout.addWidget(self.explanation_lbl)

        self.mapping_title = QLabel(tr(T.batch_columns_title))
        self.mapping_title.setStyleSheet(
            f"color: {C_DARK_GREEN}; background: transparent; font-weight: bold;"
        )
        body_layout.addWidget(self.mapping_title)

        self.mapping_container = QWidget()
        self.mapping_container.setStyleSheet(
            f"background: #FFFFFF; border: 1px solid {C_BORDER}; border-radius: 8px;"
        )
        self.mapping_layout = QVBoxLayout(self.mapping_container)
        self.mapping_layout.setContentsMargins(10, 8, 10, 8)
        self.mapping_layout.setSpacing(4)
        body_layout.addWidget(self.mapping_container)

        self.header_row_check = QCheckBox(tr(T.batch_header_row))
        self.header_row_check.setCursor(Qt.PointingHandCursor)
        self.header_row_check.setStyleSheet(
            f"QCheckBox {{ color: {C_DARK_GREEN}; background: transparent; spacing: 8px; }}"
            f"QCheckBox::indicator {{ width: 16px; height: 16px; }}"
        )
        body_layout.addWidget(self.header_row_check)

        self.format_title = QLabel(tr(T.batch_format_title))
        self.format_title.setStyleSheet(
            f"color: {C_DARK_GREEN}; background: transparent; font-weight: bold;"
        )
        body_layout.addWidget(self.format_title)

        self.options_container = QWidget()
        self.options_container.setStyleSheet("background: transparent;")
        self.options_layout = QVBoxLayout(self.options_container)
        self.options_layout.setContentsMargins(0, 0, 0, 0)
        self.options_layout.setSpacing(6)
        body_layout.addWidget(self.options_container)

        self.excel_zone = ExcelDropZone()
        body_layout.addWidget(self.excel_zone)
        self.scroll.setWidget(body)
        root.addWidget(self.scroll, 1)

        footer = QWidget()
        footer.setStyleSheet("background: transparent;")
        footer_layout = QHBoxLayout(footer)
        footer_layout.setContentsMargins(40, 8, 40, 18)
        self.confirm_btn = QPushButton(tr(T.batch_confirm))
        self.confirm_btn.setCursor(Qt.PointingHandCursor)
        self.confirm_btn.setFixedHeight(44)
        self.confirm_btn.setMinimumWidth(200)
        self.confirm_btn.setStyleSheet(f"""
            QPushButton {{
                background-color: {C_GREEN}; color: white;
                border: none; border-radius: 10px;
                font-size: 12px; font-weight: bold; padding: 0 24px;
            }}
            QPushButton:hover {{ background-color: {C_DARK_GREEN}; }}
            QPushButton:disabled {{ background-color: {C_MUTED}; }}
        """)
        self.confirm_btn.clicked.connect(self._on_confirm)
        footer_layout.addWidget(self.confirm_btn)
        footer_layout.addStretch()
        root.addWidget(footer)
        self.confirm_btn.setEnabled(False)

    def _clear_layout(self, layout):
        while layout.count():
            item = layout.takeAt(0)
            widget = item.widget()
            if widget is not None:
                widget.deleteLater()

    def load_template(self, template):
        self._template = template or {}
        self.variable_names = []
        self._font_options = {}
        self.header_row_check.setChecked(False)
        self.excel_zone.reset()
        self._clear_layout(self.mapping_layout)
        self._clear_layout(self.options_layout)

        source = Path(self._template.get("address", ""))
        if not source.is_file():
            ThemedDialog.warning(
                self, tr(T.error), tr(T.err_template_not_found, path=source)
            )
            self.confirm_btn.setEnabled(False)
            return

        try:
            properties = extract_placeholders(str(source))
        except WrapperCountError as exc:
            ThemedDialog.error(self, tr(T.err_wrapper_count), str(exc))
            self.confirm_btn.setEnabled(False)
            return
        except Exception as exc:
            ThemedDialog.error(self, tr(T.error), tr(T.cli_read_error, err=exc))
            self.confirm_btn.setEnabled(False)
            return

        self.variable_names = list(properties.keys())
        if not self.variable_names:
            ThemedDialog.warning(
                self, tr(T.error), tr(T.cli_no_variables, w=_get_wrapper())
            )
            self.confirm_btn.setEnabled(False)
            return

        self.mapping_layout.addWidget(QLabel(
            tr(T.batch_mapping_filename, column=1)
        ))
        for index, variable in enumerate(self.variable_names, start=2):
            label = QLabel(tr(
                T.batch_mapping_variable,
                column=index,
                name=variable,
            ))
            label.setWordWrap(True)
            self.mapping_layout.addWidget(label)

        for index, variable in enumerate(self.variable_names, start=2):
            card = QFrame()
            card.setStyleSheet(f"""
                QFrame {{
                    background: #FFFFFF; border: 1px solid {C_BORDER};
                    border-radius: 9px;
                }}
            """)
            row = QHBoxLayout(card)
            row.setContentsMargins(10, 6, 10, 6)
            row.setSpacing(8)

            variable_label = QLabel(tr(
                T.batch_mapping_variable,
                column=index,
                name=variable,
            ))
            variable_label.setWordWrap(True)
            variable_label.setStyleSheet(
                f"color: {C_DARK_GREEN}; background: transparent; border: none;"
            )
            row.addWidget(variable_label, 1)

            size_combo = FontSizeCombo()
            size_combo.setFixedWidth(94)
            size_combo.setToolTip(tr(T.font_size))
            row.addWidget(size_combo)

            bold_toggle = RoundToggle("B", tr(T.bold), bold=True)
            row.addWidget(bold_toggle)
            italic_toggle = RoundToggle("I", tr(T.italic), italic=True)
            row.addWidget(italic_toggle)

            color_btn = QPushButton()
            color_btn.setFixedSize(20, 20)
            color_btn.setCursor(Qt.PointingHandCursor)
            color_btn.setToolTip(tr(T.color_tooltip))
            color_btn.setProperty("color", "#000000")
            apply_color_button_style(color_btn, "#000000")
            color_btn.clicked.connect(
                lambda _=False, key=variable: self._pick_color(key)
            )
            row.addWidget(color_btn)

            self.options_layout.addWidget(card)
            self._font_options[variable] = {
                "size": size_combo,
                "bold": bold_toggle,
                "italic": italic_toggle,
                "color": color_btn,
            }

        self.confirm_btn.setEnabled(True)

    def _pick_color(self, key):
        controls = self._font_options.get(key)
        if not controls:
            return
        color_btn = controls["color"]
        current = color_btn.property("color") or "#000000"
        dialog = QColorDialog(QColor(current), self)
        dialog.setOption(QColorDialog.ColorDialogOption.DontUseNativeDialog, True)
        dialog.setOption(QColorDialog.ColorDialogOption.ShowAlphaChannel, False)
        dialog.setLayoutDirection(Qt.LeftToRight)
        dialog.setWindowTitle(tr(T.pick_color))

        palette = dialog.palette()
        palette.setColor(QPalette.ColorRole.Window, QColor("#FFFFFF"))
        palette.setColor(QPalette.ColorRole.WindowText, QColor(C_DARK_GREEN))
        palette.setColor(QPalette.ColorRole.Base, QColor("#FFFFFF"))
        palette.setColor(QPalette.ColorRole.Text, QColor(C_DARK_GREEN))
        palette.setColor(QPalette.ColorRole.Button, QColor(C_LIGHT_GREEN))
        palette.setColor(QPalette.ColorRole.ButtonText, QColor(C_DARK_GREEN))
        palette.setColor(QPalette.ColorRole.Highlight, QColor(C_GREEN))
        palette.setColor(QPalette.ColorRole.HighlightedText, QColor("#FFFFFF"))
        dialog.setPalette(palette)
        dialog.setStyleSheet(f"""
            QColorDialog {{ background: #FFFFFF; color: {C_DARK_GREEN}; }}
            QColorDialog QLabel {{ color: {C_DARK_GREEN}; background: transparent; }}
            QColorDialog QPushButton {{
                color: {C_DARK_GREEN}; background: {C_LIGHT_GREEN};
                border: 1px solid {C_BORDER}; border-radius: 6px; padding: 5px 12px;
            }}
        """)

        def localize_hex_label():
            dialog.setLayoutDirection(Qt.LeftToRight)
            for child in dialog.findChildren(QWidget):
                child.setLayoutDirection(Qt.LeftToRight)
            for label in dialog.findChildren(QLabel):
                text = label.text().replace("&", "").strip().rstrip(":").strip().casefold()
                if text == "html":
                    label.setText("HEX:")
                    break

        QTimer.singleShot(0, localize_hex_label)
        if dialog.exec():
            selected = dialog.currentColor()
            if selected.isValid():
                color_hex = selected.name()
                color_btn.setProperty("color", color_hex)
                apply_color_button_style(color_btn, color_hex)

    def _collect_font_overrides(self):
        result = {}
        for variable, controls in self._font_options.items():
            result[variable] = (
                _font_size_value(controls["size"]),
                controls["bold"].isChecked(),
                controls["italic"].isChecked(),
                controls["color"].property("color") or "#000000",
            )
        return result

    def _on_confirm(self):
        if not HAS_DOCX:
            ThemedDialog.error(self, tr(T.missing_library), tr(T.no_docx))
            return
        if not HAS_EXCEL:
            ThemedDialog.error(
                self, tr(T.missing_library), tr(T.batch_excel_library_missing)
            )
            return
        if not self._template or not self.variable_names:
            return

        excel_path = self.excel_zone.file_path()
        if not excel_path:
            ThemedDialog.warning(self, tr(T.error), tr(T.batch_excel_missing))
            return

        source = Path(self._template.get("address", ""))
        if not source.is_file():
            ThemedDialog.warning(
                self, tr(T.error), tr(T.err_template_not_found, path=source)
            )
            return

        try:
            records = read_excel_records(
                excel_path,
                self.variable_names,
                header_row=self.header_row_check.isChecked(),
            )
        except ExcelImportError as exc:
            ThemedDialog.error(self, tr(T.error), str(exc))
            return

        configured_output = str(SETTINGS.get("output_dir") or "").strip()
        base_output = Path(
            configured_output or (Path.home() / "Desktop")
        ).expanduser()
        batch_directory = base_output / tr(T.batch_output_folder)
        try:
            batch_directory.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            ThemedDialog.error(
                self, tr(T.error), tr(T.batch_output_folder_error, err=exc)
            )
            return

        reserved_names = set()
        font_overrides = self._collect_font_overrides()
        completed = 0
        for row_number, output_name, replacements in records:
            output_file = _safe_batch_filename(
                output_name, batch_directory, reserved_names
            )
            try:
                _save_docx_atomically(
                    source,
                    batch_directory / output_file,
                    replacements,
                    font_overrides,
                    overwrite=False,
                )
                completed += 1
            except Exception as exc:
                ThemedDialog.error(
                    self,
                    tr(T.error),
                    tr(T.batch_row_write_error, row=row_number, err=exc),
                )
                return

        ThemedDialog.info(
            self,
            tr(T.issued),
            tr(T.batch_success, count=completed, path=batch_directory),
        )


# ═════════════════════════════════════════════════════════════════
#  SettingsPage — no theme option
# ═════════════════════════════════════════════════════════════════

class SettingsPage(QWidget):
    back_requested = Signal()
    saved = Signal()

    IN_H      = 36
    FIELD_W   = 380
    BROWSE_W  = 80

    LANGUAGES = [
        ("fa", "فارسی"),
        ("en", "English"),
    ]

    def __init__(self):
        super().__init__()
        self.setStyleSheet("background: transparent;")

        lay = QVBoxLayout(self)
        lay.setContentsMargins(40, 22, 40, 30); lay.setSpacing(10)

        header = QHBoxLayout()
        title = QLabel(tr(T.settings))
        tf = QFont(); tf.setPointSize(15); tf.setBold(True)
        title.setFont(tf)
        title.setStyleSheet(f"color: {C_DARK_GREEN}; background: transparent;")

        back_btn = QPushButton(tr(T.back_spaced))
        back_btn.setIcon(icon_arrow_left(16)); back_btn.setIconSize(QSize(16, 16))
        back_btn.setCursor(Qt.PointingHandCursor); back_btn.setFixedHeight(34)
        back_btn.setStyleSheet(f"""
            QPushButton {{
                background: transparent; border: 1px solid {C_BORDER};
                border-radius: 8px; color: {C_DARK_GREEN};
                padding: 0 14px; font-size: 11px;
            }}
            QPushButton:hover {{
                background-color: {C_LIGHT_GREEN};
                border-color: {C_GREEN};
            }}
        """)
        back_btn.clicked.connect(self.back_requested.emit)

        header.addWidget(title); header.addStretch(); header.addWidget(back_btn)
        lay.addLayout(header)

        lf = QFont(); lf.setPointSize(10); lf.setBold(True)

        # Language
        lang_lbl = QLabel(tr(T.app_language))
        lang_lbl.setFont(lf)
        lang_lbl.setStyleSheet(f"color: {C_DARK_GREEN}; background: transparent;")
        lay.addWidget(lang_lbl)

        self.lang_combo = ComboWithArrow(editable=False)
        self.lang_combo.setFixedHeight(self.IN_H)
        self.lang_combo.setFixedWidth(self.FIELD_W)
        for code, label in self.LANGUAGES:
            self.lang_combo.addItem(label, code)
        cur_lang = SETTINGS.get("language", "fa")
        idx_lang = self.lang_combo.findData(cur_lang)
        if idx_lang >= 0:
            self.lang_combo.setCurrentIndex(idx_lang)
        self._style_combo(self.lang_combo)

        lang_row = QHBoxLayout()
        lang_row.setContentsMargins(0, 0, 0, 0)
        lang_row.addWidget(self.lang_combo)
        lang_row.addStretch()
        lay.addLayout(lang_row); lay.addSpacing(2)

        # Output dir
        out_lbl = QLabel(tr(T.output_path))
        out_lbl.setFont(lf)
        out_lbl.setStyleSheet(f"color: {C_DARK_GREEN}; background: transparent;")
        lay.addWidget(out_lbl)

        out_container = QWidget()
        out_container.setStyleSheet("background: transparent;")
        out_container.setFixedWidth(self.FIELD_W)
        out_row = QHBoxLayout(out_container)
        out_row.setContentsMargins(0, 0, 0, 0); out_row.setSpacing(8)

        self.output_input = QLineEdit()
        self.output_input.setText(SETTINGS.get("output_dir", str(Path.home() / "Desktop")))
        self.output_input.setFixedHeight(self.IN_H)
        self.output_input.setStyleSheet(f"""
            QLineEdit {{
                background-color: #FFFFFF;
                border: 1.5px solid {C_BORDER};
                border-radius: 8px;
                padding: 0 12px;
                color: {C_DARK_GREEN};
                font-size: 11px;
            }}
            QLineEdit:focus {{ border-color: {C_GREEN}; }}
        """)
        out_row.addWidget(self.output_input, 1)

        browse_btn = QPushButton(tr(T.browse))
        browse_btn.setFixedHeight(self.IN_H)
        browse_btn.setFixedWidth(self.BROWSE_W)
        browse_btn.setCursor(Qt.PointingHandCursor)
        browse_btn.setStyleSheet(f"""
            QPushButton {{
                background-color: {C_LIGHT_GREEN};
                border: 1px solid {C_BORDER};
                border-radius: 8px;
                color: {C_DARK_GREEN};
                font-size: 11px;
            }}
            QPushButton:hover {{
                background-color: {C_HOVER};
                border-color: {C_GREEN};
            }}
        """)
        browse_btn.clicked.connect(self._pick_output_dir)
        out_row.addWidget(browse_btn)

        out_wrap = QHBoxLayout()
        out_wrap.setContentsMargins(0, 0, 0, 0)
        out_wrap.addWidget(out_container)
        out_wrap.addStretch()
        lay.addLayout(out_wrap); lay.addSpacing(2)

        # Wrapper
        wrap_lbl = QLabel(tr(T.placeholder_wrapper))
        wrap_lbl.setFont(lf)
        wrap_lbl.setStyleSheet(f"color: {C_DARK_GREEN}; background: transparent;")
        lay.addWidget(wrap_lbl)

        self.wrapper_input = QLineEdit()
        self.wrapper_input.setText(SETTINGS.get("placeholder_wrapper", '##'))
        self.wrapper_input.setFixedHeight(self.IN_H)
        self.wrapper_input.setFixedWidth(self.FIELD_W)
        self.wrapper_input.setPlaceholderText(tr(T.ph_wrapper))
        self.wrapper_input.setStyleSheet(f"""
            QLineEdit {{
                background-color: #FFFFFF;
                border: 1.5px solid {C_BORDER};
                border-radius: 8px;
                padding: 0 12px;
                color: {C_DARK_GREEN};
                font-size: 12px;
                font-family: Consolas, monospace;
            }}
            QLineEdit:focus {{ border-color: {C_GREEN}; }}
        """)

        wrap_row = QHBoxLayout()
        wrap_row.setContentsMargins(0, 0, 0, 0)
        wrap_row.addWidget(self.wrapper_input)
        wrap_row.addStretch()
        lay.addLayout(wrap_row)

        example = self.wrapper_input.text() + tr(T.user_name) + self.wrapper_input.text()
        hint = QLabel(tr(T.wrapper_hint, example=example))
        hint.setStyleSheet(f"color: {C_MUTED}; background: transparent; font-size: 10px;")
        hint.setWordWrap(True)
        lay.addWidget(hint)

        self.launch_startup_check = QCheckBox(tr(T.launch_on_startup))
        self.launch_startup_check.setChecked(
            SETTINGS.get("launch_on_startup") is True
        )
        self.launch_startup_check.setToolTip(tr(T.launch_on_startup_tooltip))
        self.launch_startup_check.setLayoutDirection(_dir())
        self.launch_startup_check.setStyleSheet(f"""
            QCheckBox {{
                color: {C_DARK_GREEN}; background: transparent;
                font-size: 11px; spacing: 8px;
            }}
            QCheckBox::indicator {{
                width: 18px; height: 18px; border: 1px solid {C_BORDER};
                border-radius: 4px; background: #FFFFFF;
            }}
            QCheckBox::indicator:checked {{
                background: {C_GREEN}; border-color: {C_GREEN};
            }}
        """)
        if sys.platform == "win32":
            lay.addWidget(self.launch_startup_check)

        lay.addStretch()

        btn_row = QHBoxLayout()
        self.save_btn = QPushButton(tr(T.save_settings))
        self.save_btn.setCursor(Qt.PointingHandCursor)
        self.save_btn.setFixedHeight(42); self.save_btn.setMinimumWidth(160)
        self.save_btn.setStyleSheet(f"""
            QPushButton {{
                background-color: {C_GREEN}; color: white;
                border: none; border-radius: 10px;
                font-size: 12px; font-weight: bold;
                padding: 0 24px;
            }}
            QPushButton:hover {{ background-color: {C_DARK_GREEN}; }}
        """)
        self.save_btn.clicked.connect(self._on_save)
        btn_row.addWidget(self.save_btn)
        btn_row.addStretch()
        lay.addLayout(btn_row)

    def _style_combo(self, combo):
        combo.setStyleSheet(f"""
            QComboBox {{
                background-color: #FFFFFF;
                border: 1.5px solid {C_BORDER};
                border-radius: 8px;
                padding-left: 26px;
                padding-right: 10px;
                color: {C_DARK_GREEN};
                font-size: 11px;
            }}
            QComboBox:focus, QComboBox:on {{ border-color: {C_GREEN}; }}
            QComboBox QAbstractItemView {{
                background-color: #FFFFFF;
                border: 1px solid {C_BORDER};
                border-radius: 8px;
                padding: 4px;
                outline: none;
                color: {C_DARK_GREEN};
                selection-background-color: {C_LIGHT_GREEN};
                selection-color: {C_DARK_GREEN};
            }}
            QComboBox QAbstractItemView::item {{
                padding: 6px; border-radius: 6px; min-height: 20px;
            }}
            QComboBox QAbstractItemView::item:hover {{
                background-color: {C_HOVER};
            }}
        """)

    def _pick_output_dir(self):
        start = self.output_input.text() or str(Path.home())
        d = QFileDialog.getExistingDirectory(self, tr(T.select_output_dir), start)
        if d:
            self.output_input.setText(d)

    def _on_save(self):
        new_lang = self.lang_combo.currentData()
        new_output = self.output_input.text().strip()
        new_wrapper = self.wrapper_input.text()

        if not new_output:
            ThemedDialog.warning(self, tr(T.error), tr(T.err_output_empty))
            return
        try:
            Path(new_output).mkdir(parents=True, exist_ok=True)
        except Exception as e:
            ThemedDialog.error(self, tr(T.error),
                               tr(T.err_output_invalid, err=e))
            return

        if len(new_wrapper) == 0:
            ThemedDialog.warning(self, tr(T.error), tr(T.err_wrapper_empty))
            return

        new_launch_on_startup = False
        if sys.platform == "win32":
            new_launch_on_startup = self.launch_startup_check.isChecked()
            try:
                _set_windows_startup(new_launch_on_startup)
            except Exception as exc:
                ThemedDialog.error(
                    self,
                    tr(T.error),
                    tr(T.startup_update_error, err=exc),
                )
                return

        SETTINGS["language"] = new_lang
        SETTINGS["output_dir"] = new_output
        SETTINGS["placeholder_wrapper"] = new_wrapper
        SETTINGS["launch_on_startup"] = new_launch_on_startup
        save_settings()

        self.saved.emit()


# ═════════════════════════════════════════════════════════════════
#  Container — plain background, no decorations
# ═════════════════════════════════════════════════════════════════

class Container(QWidget):
    def __init__(self):
        super().__init__()
        self.setStyleSheet("background: transparent;")

    def paintEvent(self, e):
        p = QPainter(self)
        p.setRenderHint(QPainter.Antialiasing)
        w, h = self.width(), self.height()

        try:
            is_max = self.window().isMaximized()
        except Exception:
            is_max = False
        radius = 0 if is_max else 14

        path = QPainterPath()
        path.addRoundedRect(QRectF(0.5, 0.5, w - 1, h - 1), radius, radius)
        p.fillPath(path, QColor(C_BG))

        if radius > 0:
            p.setPen(QPen(QColor(C_BORDER), 1))
            p.setBrush(Qt.NoBrush)
            p.drawPath(path)


# ═════════════════════════════════════════════════════════════════
#  MainWindow
# ═════════════════════════════════════════════════════════════════

class MainWindow(QMainWindow):
    theme_changed = Signal()
    RESIZE_MARGIN = 6

    def __init__(self):
        super().__init__()
        self._built_language = SETTINGS.get("language", "fa")
        self._really_quit = False
        self.setWindowTitle(tr(T.app_title))
        self.setWindowFlags(Qt.FramelessWindowHint | Qt.Window)
        self.setAttribute(Qt.WA_TranslucentBackground)
        saved_size = SETTINGS.get("window_size", [720, 600])
        try:
            if not isinstance(saved_size, (list, tuple)) or len(saved_size) < 2:
                raise ValueError("invalid saved window size")
            start_w = max(300, int(saved_size[0]))
            start_h = max(300, int(saved_size[1]))
        except (TypeError, ValueError, IndexError):
            start_w, start_h = 720, 600
        screen = QGuiApplication.primaryScreen()
        if screen is not None:
            available = screen.availableGeometry()
            start_w = min(start_w, max(300, available.width()))
            start_h = min(start_h, max(300, available.height()))
        self.resize(start_w, start_h)
        self._window_state_timer = QTimer(self)
        self._window_state_timer.setSingleShot(True)
        self._window_state_timer.timeout.connect(self._save_window_state)
        self.setMinimumSize(300, 300)

        self.setWindowFlags(Qt.Window | Qt.FramelessWindowHint)

        root = QWidget()
        root.setAttribute(Qt.WA_TranslucentBackground)
        self._root_layout = QVBoxLayout(root)
        self._root_layout.setContentsMargins(0, 0, 0, 0)

        self.container = Container()
        self._root_layout.addWidget(self.container)
        self.setCentralWidget(root)

        v = QVBoxLayout(self.container)
        v.setContentsMargins(0, 0, 0, 0); v.setSpacing(0)
        self.title_bar = TitleBar(self)
        v.addWidget(self.title_bar)

        self.stack = QStackedWidget()
        self.stack.setStyleSheet("background: transparent;")
        v.addWidget(self.stack, 1)

        self.home_page = HomePage()
        self.add_page = AddLetterPage()
        self.edit_page = EditLetterPage()
        self.bulk_page = BulkIssuePage()
        self.settings_page = SettingsPage()

        for w in (
            self.home_page, self.add_page, self.edit_page,
            self.bulk_page, self.settings_page,
        ):
            self.stack.addWidget(w)

        self.home_page.dropdown.add_requested.connect(self._go_to_add)
        self.home_page.dropdown.template_selected.connect(self._on_template_selected)
        self.home_page.dropdown.template_deleted.connect(self._on_template_deleted)
        self.home_page.settings_requested.connect(self._go_to_settings)
        self.home_page.quick_template_selected.connect(self._on_template_selected)
        self.home_page.quick_access_changed.connect(self._refresh_tray_menu)
        
        self.add_page.back_requested.connect(self._go_to_home)
        self.add_page.saved.connect(self._on_letter_saved)

        self.edit_page.back_requested.connect(self._go_to_home)
        self.edit_page.batch_requested.connect(self._open_bulk_issue)
        self.bulk_page.back_requested.connect(self._return_to_issue_page)

        self.settings_page.back_requested.connect(self._go_to_home)
        self.settings_page.saved.connect(self._on_settings_saved)

        self.reload_templates()
        self.reload_templates()
        self._setup_tray()

    # ═════════════════════════════════════════════════════════════
    #  System Tray
    # ═════════════════════════════════════════════════════════════

    def _setup_tray(self):
        self.tray_icon = QSystemTrayIcon(self)
        icon_path = APP_ROOT / "app.ico"
        if icon_path.exists():
            self.tray_icon.setIcon(QIcon(str(icon_path)))
        else:
            self.tray_icon.setIcon(icon_document(32, C_DARK_GREEN))
        self.tray_icon.setToolTip(tr(T.app_title))

        self.tray_menu = QMenu()
        self._build_tray_menu()

        self.tray_icon.setContextMenu(self.tray_menu)
        self.tray_icon.activated.connect(self._on_tray_activated)
        self.tray_icon.show()

    def _build_tray_menu(self):
        self.tray_menu.clear()

        # ── باز کردن برنامه ──
        open_action = QAction(
            tr({"persian": "باز کردن برنامه", "english": "Open App"}), self)
        open_action.triggered.connect(self._show_from_tray)
        self.tray_menu.addAction(open_action)

        # ── قالب‌های Quick Access ──
        qa_names = [n for n in SETTINGS.get("quick_access", []) if n]
        if qa_names:
            templates = load_templates()
            tpl_map = {t.get("name", ""): t for t in templates}
            valid_qa = [n for n in qa_names if n in tpl_map]

            if valid_qa:
                self.tray_menu.addSeparator()

                header = QAction(tr(T.quick_access), self)
                header.setEnabled(False)
                self.tray_menu.addAction(header)

                for name in valid_qa:
                    act = QAction(name, self)
                    act.triggered.connect(
                        lambda _=None, n=name:
                            self._open_template_from_tray(n))
                    self.tray_menu.addAction(act)

        # ── خروج ──
        self.tray_menu.addSeparator()
        exit_action = QAction(
            tr({"persian": "خروج", "english": "Exit"}), self)
        exit_action.triggered.connect(self._quit_app)
        self.tray_menu.addAction(exit_action)

    def _refresh_tray_menu(self):
        if hasattr(self, "tray_menu"):
            self._build_tray_menu()

    def _show_from_tray(self):
        self.show()
        if self.isMinimized():
            self.showNormal()
        self.raise_()
        self.activateWindow()

    def _open_template_from_tray(self, name):
        templates = load_templates()
        tpl_map = {t.get("name", ""): t for t in templates}
        tpl = tpl_map.get(name)
        if tpl is None:
            self._refresh_tray_menu()
            return
        self._show_from_tray()
        self._on_template_selected(tpl)

    def _on_tray_activated(self, reason):
        if reason in (QSystemTrayIcon.Trigger, QSystemTrayIcon.DoubleClick):
            if self.isVisible() and not self.isMinimized():
                # کلیک روی آیکن وقتی پنجره باز است → پنهان کن
                self.hide()
            else:
                self._show_from_tray()

    def _quit_app(self):
        self._really_quit = True
        # مستقیم از طریق QApplication خارج شو تا از closeEvent رد شویم.
        QApplication.instance().quit()

    def _apply_win32_styles(self):
        """استایل WS_THICKFRAME و WS_CAPTION را در سطح Win32 فعال می‌کند
        تا ویندوز پنجره را به عنوان یک پنجره قابل Snap و Resize بشناسد."""
        if sys.platform != "win32":
            return
        try:
            from win32gui import GetWindowLong, SetWindowLong
            from win32con import WS_THICKFRAME, WS_CAPTION, WS_SYSMENU, WS_MINIMIZEBOX, WS_MAXIMIZEBOX, GWL_STYLE

            hwnd = int(self.winId())

            style = GetWindowLong(hwnd, GWL_STYLE)
            style |= WS_THICKFRAME      # برای Resize و Snap
            style |= WS_CAPTION         # برای Drag و Snap
            style |= WS_SYSMENU         # برای منوی سیستم
            style |= WS_MINIMIZEBOX     # برای مینیمایز
            style |= WS_MAXIMIZEBOX     # برای مکسیمایز
            SetWindowLong(hwnd, GWL_STYLE, style)
        except Exception as e:
            print(f"[Aero Snap] init styles failed: {e}")

    def nativeEvent(self, eventType, message):
        if sys.platform == "win32":
            try:
                import ctypes
                from ctypes import wintypes
                try:
                    addr = int(message)
                except Exception:
                    addr = ctypes.cast(message, ctypes.c_void_p).value
                msg = wintypes.MSG.from_address(addr)

                WM_NCCALCSIZE = 0x0083
                WM_NCHITTEST = 0x0084

                # ─── حذف حاشیه‌های Non-Client تا عنوان بومی حذف شود،
                #     اما استایل WS_THICKFRAME باقی بماند و Snap کار کند ───
                if msg.message == WM_NCCALCSIZE:
                    # مقدار 0 برگردانده می‌شود تا کل پنجره «Client Area» در نظر گرفته شود
                    return True, 0

                # ─── تعیین نواحی Resize روی لبه‌ها ───
                if msg.message == WM_NCHITTEST and not self.isMaximized():
                    x = ctypes.c_short(msg.lParam & 0xFFFF).value
                    y = ctypes.c_short((msg.lParam >> 16) & 0xFFFF).value

                    dpr = 1.0
                    try:
                        wh = self.windowHandle()
                        if wh is not None:
                            dpr = wh.devicePixelRatio() or 1.0
                    except Exception:
                        pass

                    local = self.mapFromGlobal(QPoint(int(x / dpr), int(y / dpr)))
                    w, h = self.width(), self.height()
                    m = self.RESIZE_MARGIN

                    on_l = 0 <= local.x() < m
                    on_r = w - m < local.x() <= w
                    on_t = 0 <= local.y() < m
                    on_b = h - m < local.y() <= h

                    if on_t and on_l: return True, 13  # HTTOPLEFT
                    if on_t and on_r: return True, 14  # HTTOPRIGHT
                    if on_b and on_l: return True, 16  # HTBOTTOMLEFT
                    if on_b and on_r: return True, 17  # HTBOTTOMRIGHT
                    if on_l:          return True, 10  # HTLEFT
                    if on_r:          return True, 11  # HTRIGHT
                    if on_t:          return True, 12  # HTTOP
                    if on_b:          return True, 15  # HTBOTTOM

            except Exception:
                pass

        return super().nativeEvent(eventType, message)

    # ─── فراخوانی _apply_win32_styles پس از نمایش پنجره ───
    def showEvent(self, e):
        super().showEvent(e)
        self._apply_win32_styles()
        
    def reload_templates(self):
        self.home_page.set_templates(load_templates())
        if hasattr(self, "tray_menu"):
            self._refresh_tray_menu()

    def _go_to_home(self):
        self.stack.setCurrentWidget(self.home_page)

    def _open_bulk_issue(self):
        template = self.edit_page._template
        if not template:
            return
        self.bulk_page.load_template(template)
        self.stack.setCurrentWidget(self.bulk_page)

    def _return_to_issue_page(self):
        self.stack.setCurrentWidget(self.edit_page)

    def _go_to_add(self):
        self.add_page.reset()
        self.stack.setCurrentWidget(self.add_page)

    def _go_to_settings(self):
        self.stack.removeWidget(self.settings_page)
        self.settings_page.deleteLater()
        self.settings_page = SettingsPage()
        self.settings_page.back_requested.connect(self._go_to_home)
        self.settings_page.saved.connect(self._on_settings_saved)
        self.stack.addWidget(self.settings_page)
        self.stack.setCurrentWidget(self.settings_page)

    def _on_settings_saved(self):
        self._refresh_tray_menu()
        lang_changed = SETTINGS.get("language", "fa") != self._built_language
        if lang_changed:
            self.theme_changed.emit()
        else:
            self._go_to_home()

    def _on_letter_saved(self):
        self.reload_templates()
        self._go_to_home()

    def _on_template_selected(self, tpl):
        if not HAS_DOCX:
            ThemedDialog.error(self, tr(T.missing_library), tr(T.no_docx))
            return
        self.edit_page.load_template(tpl)
        self.stack.setCurrentWidget(self.edit_page)

    def _on_template_deleted(self, tpl):
        name = tpl.get("name", "")
        address = tpl.get("address", "")

        ok = ThemedDialog.confirm_delete(
            self, tr(T.delete_template),
            tr(T.confirm_delete_tpl, name=name)
        )
        if not ok: return

        data = load_templates_raw()
        others = [t for t in data
                  if t.get("name") != name and t.get("address") == address]

        data = [t for t in data if t.get("name") != name]
        save_templates(data)

        if not others and address:
            p = Path(address)
            if p.exists():
                try:
                    p.unlink()
                except Exception as e:
                    print(f"[error] file removal: {e}")

        qa = [n for n in SETTINGS.get("quick_access", []) if n != name]
        if qa != SETTINGS.get("quick_access", []):
            SETTINGS["quick_access"] = qa
            save_settings()

        self.reload_templates()

        ThemedDialog.info(self, tr(T.deleted),
                          tr(T.deleted_ok, name=name))

    def _schedule_window_state_save(self):
        timer = getattr(self, "_window_state_timer", None)
        if timer is not None:
            timer.start(350)

    def _save_window_state(self):
        try:
            if self.isMinimized():
                save_settings()
                return
            if self.isMaximized():
                SETTINGS["window_maximized"] = True
                normal = self.normalGeometry()
                if normal.isValid():
                    SETTINGS["window_size"] = [
                        max(300, int(normal.width())),
                        max(300, int(normal.height())),
                    ]
            else:
                SETTINGS["window_maximized"] = False
                SETTINGS["window_size"] = [
                    max(300, int(self.width())),
                    max(300, int(self.height())),
                ]
            save_settings()
        except Exception as exc:
            print(f"[window] Could not save window state: {exc}")

    def resizeEvent(self, event):
        super().resizeEvent(event)
        self._schedule_window_state_save()

    def closeEvent(self, event):
        timer = getattr(self, "_window_state_timer", None)
        if timer is not None:
            timer.stop()
        self._save_window_state()

        if not getattr(self, "_really_quit", False):
            # ── فقط مخفی کن، برنامه در tray باقی می‌ماند ──
            event.ignore()
            self.hide()
            return

        # ── خروج واقعی ──
        if hasattr(self, "tray_icon"):
            self.tray_icon.hide()
        super().closeEvent(event)

    def changeEvent(self, e):
        if e.type() == QEvent.WindowStateChange:
            try:
                self.title_bar.refresh_radius()
            except Exception:
                pass
            try:
                self.container.update()
            except Exception:
                pass
            if self.isMaximized():
                SETTINGS["window_maximized"] = True
            elif not self.isMinimized():
                SETTINGS["window_maximized"] = False
            self._schedule_window_state_save()
        super().changeEvent(e)


# ═════════════════════════════════════════════════════════════════
#  First-run Language Selection
# ═════════════════════════════════════════════════════════════════

class LangButton(QFrame):
    """دکمه‌ی بزرگ انتخاب زبان با افکت hover."""
    clicked = Signal()

    def __init__(self, primary, secondary, parent=None):
        super().__init__(parent)
        self.setObjectName("langbtn")
        self.setCursor(Qt.PointingHandCursor)
        self.setMinimumHeight(160)

        lay = QVBoxLayout(self)
        lay.setContentsMargins(20, 20, 20, 20)
        lay.setSpacing(4)
        lay.addStretch()

        self._primary = QLabel(primary)
        pf = QFont(); pf.setPointSize(26); pf.setBold(True)
        self._primary.setFont(pf)
        self._primary.setAlignment(Qt.AlignCenter)
        self._primary.setAttribute(Qt.WA_TransparentForMouseEvents)

        self._secondary = QLabel(secondary)
        sf = QFont(); sf.setPointSize(11)
        self._secondary.setFont(sf)
        self._secondary.setAlignment(Qt.AlignCenter)
        self._secondary.setAttribute(Qt.WA_TransparentForMouseEvents)

        lay.addWidget(self._primary)
        lay.addWidget(self._secondary)
        lay.addStretch()

        self._apply_style(hover=False)

    def _apply_style(self, hover):
        if hover:
            self.setStyleSheet(f"""
                QFrame#langbtn {{
                    background-color: {C_GREEN};
                    border: 2px solid {C_DARK_GREEN};
                    border-radius: 16px;
                }}
            """)
            self._primary.setStyleSheet("color: white; background: transparent;")
            self._secondary.setStyleSheet(
                "color: #E6EEE8; background: transparent;")
        else:
            self.setStyleSheet(f"""
                QFrame#langbtn {{
                    background-color: {C_LIGHT_GREEN};
                    border: 2px solid {C_BORDER};
                    border-radius: 16px;
                }}
            """)
            self._primary.setStyleSheet(
                f"color: {C_DARK_GREEN}; background: transparent;")
            self._secondary.setStyleSheet(
                f"color: {C_MUTED}; background: transparent;")

    def enterEvent(self, e):
        self._apply_style(True)
        super().enterEvent(e)

    def leaveEvent(self, e):
        self._apply_style(False)
        super().leaveEvent(e)

    def mousePressEvent(self, e):
        if e.button() == Qt.LeftButton:
            self.clicked.emit()


class LanguageSelectDialog(QDialog):
    """دیالوگ انتخاب زبان که فقط در اولین اجرای برنامه نمایش داده می‌شود."""

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowFlags(
            Qt.Dialog | Qt.FramelessWindowHint | Qt.WindowStaysOnTopHint)
        self.setAttribute(Qt.WA_TranslucentBackground)
        self.setModal(True)
        # مستقل از زبان برنامه، این دیالوگ همیشه LTR است تا هر دو دکمه
        # در جای متقارن دیده شوند.
        self.setLayoutDirection(Qt.LeftToRight)
        self.setFixedSize(460, 580)

        self.selected_language = None

        outer = QVBoxLayout(self)
        outer.setContentsMargins(18, 18, 18, 18)

        frame = QFrame()
        frame.setObjectName("langFrame")
        frame.setStyleSheet(f"""
            QFrame#langFrame {{
                background-color: #FFFFFF;
                border-radius: 20px;
                border: 1px solid {C_BORDER};
            }}
        """)
        sh = QGraphicsDropShadowEffect()
        sh.setBlurRadius(44)
        sh.setOffset(0, 12)
        sh.setColor(QColor(0, 0, 0, 95))
        frame.setGraphicsEffect(sh)
        outer.addWidget(frame)

        v = QVBoxLayout(frame)
        v.setContentsMargins(28, 30, 28, 28)
        v.setSpacing(14)

        title = QLabel("Choose your language  •  زبان خود را انتخاب کنید")
        tf = QFont(); tf.setPointSize(12); tf.setBold(True)
        title.setFont(tf)
        title.setAlignment(Qt.AlignCenter)
        title.setWordWrap(True)
        title.setStyleSheet(
            f"color: {C_DARK_GREEN}; background: transparent;")
        v.addWidget(title)

        v.addSpacing(6)

        persian_btn = LangButton("فارسی", "Persian")
        english_btn = LangButton("English", "انگلیسی")

        persian_btn.clicked.connect(lambda: self._choose("fa"))
        english_btn.clicked.connect(lambda: self._choose("en"))

        # stretch factor = 1 ⇒ هر دکمه نیمی از فضای عمودی را می‌گیرد.
        v.addWidget(persian_btn, 1)
        v.addWidget(english_btn, 1)

    def _choose(self, lang):
        self.selected_language = lang
        self.accept()

    def showEvent(self, e):
        super().showEvent(e)
        scr = QGuiApplication.primaryScreen()
        if scr is not None:
            geo = scr.availableGeometry()
            self.move(geo.center().x() - self.width() // 2,
                      geo.center().y() - self.height() // 2)

# ═════════════════════════════════════════════════════════════════
#  Command-line document issuance
# ═════════════════════════════════════════════════════════════════

def _build_cli_parser():
    """Create the parser used by both python and the frozen executable."""
    parser = argparse.ArgumentParser(
        prog=Path(sys.argv[0]).name or "WordCloner",
        description="Issue a WordCloner document from a template or Excel workbook.",
        epilog=(
            "Single document:\n"
            '  WordCloner.exe -i "letter_template.docx" -o "output.docx" '
            '-v "A. User" " " "Operations"\n\n'
            "Count template variables:\n"
            '  WordCloner.exe -i "letter_template.docx" --count\n\n'
            "Bulk from Excel:\n"
            '  WordCloner.exe -i "letter_template.docx" --excel "data.xlsx" '
            "--header-row"
        ),
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument(
        "-i", "--input", dest="input_path", required=True, metavar="DOCX",
        help="Path to the .docx template.",
    )
    parser.add_argument(
        "-o", "--output", dest="output_path", default=None, metavar="PATH",
        help=(
            "Output .docx file for single issuance; output folder for --excel "
            "(default: settings output folder / Bulk Issuance)."
        ),
    )
    parser.add_argument(
        "-v", "--values", dest="values", nargs="*", default=None,
        metavar="VALUE",
        help=(
            'Values in first-appearance order. Pass a quoted single space (" ") '
            "for an intentionally blank variable."
        ),
    )
    parser.add_argument(
        "-w", "--wrapper", dest="wrapper", default=None, metavar="TEXT",
        help="Use TEXT as the placeholder wrapper for this run (default: settings.json).",
    )
    parser.add_argument(
        "-x", "--excel", dest="excel_path", default=None, metavar="XLSX",
        help="Read one output filename and the variable values from each Excel row.",
    )
    parser.add_argument(
        "--header-row", action="store_true",
        help="Skip row 1 of the active worksheet because it contains column headings.",
    )
    parser.add_argument(
        "-c", "--count", action="store_true",
        help="Print the number of unique variables in the template and exit.",
    )
    parser.add_argument(
        "--force", action="store_true",
        help="Replace an existing single-document output file.",
    )
    return parser


def _cli_output_path(raw_path, parser):
    path = Path(raw_path).expanduser()
    if path.exists() and path.is_dir():
        parser.error(tr(T.cli_output_is_directory, path=path))
    if path.suffix.lower() != ".docx":
        path = Path(str(path) + ".docx")
    return path


def cli_main(argv=None):
    """Handle CLI operations without constructing QApplication or a GUI page."""
    parser = _build_cli_parser()
    args = parser.parse_args(argv)

    # CLI messages are always English, independently of the saved GUI language.
    # This is an in-memory presentation override; the CLI never saves it.
    SETTINGS["language"] = "en"

    if not HAS_DOCX:
        parser.error(tr(T.no_docx))

    try:
        migrate_legacy_data()
        load_settings()
        # load_settings() restores the user's GUI language, so reapply the
        # CLI-only English override after reading settings.json.
        SETTINGS["language"] = "en"
    except Exception as exc:
        SETTINGS["language"] = "en"
        parser.error(tr(T.cli_read_error, err=exc))

    if args.wrapper is not None:
        if not args.wrapper.strip():
            parser.error(tr(T.err_wrapper_empty))
        # A CLI wrapper override applies to this invocation only; it does not
        # silently rewrite the user's persistent GUI settings.
        SETTINGS["placeholder_wrapper"] = args.wrapper

    input_path = Path(args.input_path).expanduser()
    if not input_path.exists():
        parser.error(tr(T.err_template_not_found, path=input_path))
    if not input_path.is_file():
        parser.error(tr(T.cli_input_not_file, path=input_path))
    if input_path.suffix.lower() != ".docx":
        parser.error(tr(T.cli_input_not_docx, path=input_path))
    input_path = input_path.resolve()

    wrapper = _get_wrapper()
    if not isinstance(wrapper, str) or not wrapper:
        parser.error(tr(T.err_wrapper_empty))

    try:
        properties = extract_placeholders(str(input_path))
    except WrapperCountError as exc:
        parser.error(str(exc))
    except Exception as exc:
        parser.error(tr(T.cli_read_error, err=exc))
    variable_names = list(properties.keys())

    if args.count:
        if (args.output_path is not None or args.values is not None
                or args.excel_path is not None or args.header_row or args.force):
            parser.error(tr(T.cli_count_mode_conflict))
        print(tr(T.cli_variable_count, count=len(variable_names)))
        return 0

    if args.header_row and args.excel_path is None:
        parser.error(tr(T.cli_header_only_excel))

    if args.excel_path is not None:
        if args.values is not None or args.force:
            parser.error(tr(T.cli_excel_mode_conflict))
        if not variable_names:
            parser.error(tr(T.cli_no_variables, w=wrapper))
        if not HAS_EXCEL:
            parser.error(tr(T.batch_excel_library_missing))

        try:
            records = read_excel_records(
                args.excel_path, variable_names, header_row=args.header_row
            )
        except ExcelImportError as exc:
            parser.error(str(exc))

        if args.output_path:
            batch_directory = Path(args.output_path).expanduser()
            if batch_directory.suffix.lower() == ".docx":
                parser.error(tr(
                    T.cli_excel_output_folder_error, path=batch_directory
                ))
            if batch_directory.exists() and not batch_directory.is_dir():
                parser.error(tr(
                    T.cli_excel_output_folder_error, path=batch_directory
                ))
        else:
            configured_output = str(SETTINGS.get("output_dir") or "").strip()
            output_root = Path(
                configured_output or (Path.home() / "Desktop")
            ).expanduser()
            batch_directory = output_root / tr(T.batch_output_folder)

        try:
            batch_directory.mkdir(parents=True, exist_ok=True)
        except OSError as exc:
            parser.error(tr(T.batch_output_folder_error, err=exc))

        reserved_names = set()
        for row_number, output_name, replacements in records:
            output_file = _safe_batch_filename(
                output_name, batch_directory, reserved_names
            )
            try:
                _save_docx_atomically(
                    input_path,
                    batch_directory / output_file,
                    replacements,
                    {},
                    overwrite=False,
                )
            except Exception as exc:
                parser.error(tr(
                    T.batch_row_write_error, row=row_number, err=exc
                ))

        print(tr(T.batch_success, count=len(records), path=batch_directory))
        return 0

    if not variable_names:
        parser.error(tr(T.cli_no_variables, w=wrapper))
    if not args.output_path:
        parser.error(tr(T.cli_output_required))

    values = args.values or []
    if len(values) > len(variable_names):
        parser.error(tr(
            T.cli_too_many_values,
            expected=len(variable_names),
            actual=len(values),
        ))
    if len(values) < len(variable_names):
        parser.error(tr(
            T.cli_too_few_values,
            expected=len(variable_names),
            actual=len(values),
        ))
    if any(value == "" for value in values):
        parser.error(tr(T.cli_empty_value))

    replacements = dict(zip(variable_names, values))
    output_path = _cli_output_path(args.output_path, parser).resolve()
    if output_path == input_path:
        parser.error(tr(T.cli_output_is_input))
    if output_path.exists() and not args.force:
        parser.error(tr(T.cli_output_exists, path=output_path))

    try:
        _save_docx_atomically(
            input_path, output_path, replacements, {}, overwrite=args.force
        )
    except FileExistsError:
        parser.error(tr(T.cli_output_exists, path=output_path))
    except Exception as exc:
        parser.error(tr(T.cli_write_error, err=exc))

    print(tr(T.cli_success, path=output_path, count=len(variable_names)))
    return 0


def _attach_cli_console_if_needed():
    """Attach CLI streams to a parent console if this launcher lacks them.

    The recommended Windows spec is console-enabled with PyInstaller's
    hide_console='hide-late': this lets CMD wait for CLI completion while an
    owned console is hidden for GUI startup. This helper is only a fallback
    for launchers that start the CLI without standard streams; attaching a
    windowed executable cannot make its parent CMD wait synchronously.
    """
    if sys.platform != "win32" or (sys.stdout is not None and sys.stderr is not None):
        return
    try:
        import ctypes

        kernel32 = ctypes.windll.kernel32
        kernel32.AttachConsole(-1)  # ATTACH_PARENT_PROCESS
        if not kernel32.GetConsoleWindow():
            kernel32.AllocConsole()
        if sys.stdin is None:
            sys.stdin = open("CONIN$", "r", encoding="utf-8", errors="replace")
        if sys.stdout is None:
            sys.stdout = open(
                "CONOUT$", "w", encoding="utf-8", errors="replace", buffering=1
            )
        if sys.stderr is None:
            sys.stderr = open(
                "CONOUT$", "w", encoding="utf-8", errors="replace", buffering=1
            )
    except Exception:
        # Leave the normal Python streams alone when running outside the
        # expected Windows console environment.
        pass


# ═════════════════════════════════════════════════════════════════
#  GUI startup
# ═════════════════════════════════════════════════════════════════

def gui_main(start_hidden=False):
    migrate_legacy_data()

    # ─── تشخیص اولین اجرا قبل از بارگذاری settings ───
    first_run = not SETTINGS_JSON.exists()

    load_settings()
    _sync_windows_startup_from_settings()

    if sys.platform == "win32":
        myappid = "com.wordcloner.letterissuance.1.0"
        try:
            from ctypes import windll
            windll.shell32.SetCurrentProcessExplicitAppUserModelID(myappid)
        except Exception as e:
            print(f"[AppUserModelID] failed: {e}")

    # Strip any launcher-specific arguments before Qt initializes.
    app = QApplication([sys.argv[0]])
    icon_path = APP_ROOT / "app.ico"
    app.setWindowIcon(QIcon(str(icon_path)))
    app.setQuitOnLastWindowClosed(False)

    if start_hidden and not QSystemTrayIcon.isSystemTrayAvailable():
        print("[startup] System tray is unavailable; showing the main window.")
        start_hidden = False

    for fam in ("Vazirmatn", "IRANSans", "Tahoma", "Segoe UI"):
        f = QFont(fam, 10)
        if f.exactMatch() or fam == "Segoe UI":
            app.setFont(f)
            break

    # ─── اولین اجرا: انتخاب زبان؛ در شروع مخفی dialog نمایش داده نشود ───
    if first_run:
        if not start_hidden:
            dlg = LanguageSelectDialog()
            if dlg.exec() == QDialog.Accepted and dlg.selected_language:
                SETTINGS["language"] = dlg.selected_language
            # در غیر این صورت زبان پیش‌فرض (fa) باقی می‌ماند.
        save_settings()

    # setLayoutDirection پس از دیالوگ تا دیالوگ LTR خالص بماند.
    app.setLayoutDirection(_dir())

    current = [None]

    def restart():
        old = current[0]
        try:
            old._save_window_state()
            if hasattr(old, "_window_state_timer"):
                old._window_state_timer.stop()
        except Exception:
            pass
        was_max = old.isMaximized()
        geo = old.normalGeometry() if was_max else old.geometry()
        if not geo.isValid():
            geo = old.geometry()
        old.hide()
        old.deleteLater()

        load_settings()
        app.setLayoutDirection(_dir())

        new_win = MainWindow()
        new_win.theme_changed.connect(restart)
        new_win.setGeometry(geo)
        if was_max:
            new_win.showMaximized()
        else:
            new_win.show()
        current[0] = new_win

    def create_window():
        win = MainWindow()
        win.theme_changed.connect(restart)
        if not start_hidden:
            if SETTINGS.get("window_maximized") is True:
                win.showMaximized()
            else:
                win.show()
        return win

    current[0] = create_window()
    return app.exec()


def main(argv=None):
    """Dispatch to CLI when arguments are supplied; otherwise launch the GUI.

    This same entry point is used by the Python console command and the frozen
    WordCloner.exe, so opening the EXE without arguments starts the GUI.
    """
    arguments = list(sys.argv[1:] if argv is None else argv)
    if arguments == ["--gui"]:
        return gui_main()
    if arguments == ["--startup"]:
        return gui_main(start_hidden=True)
    if arguments:
        _attach_cli_console_if_needed()
        return cli_main(arguments)
    return gui_main()


if __name__ == "__main__":
    raise SystemExit(main())